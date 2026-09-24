"""文档读写 Skills。

智能体可通过这些工具生成 / 读取 Word、PPT、PDF、CSV、Excel 等文件，
产物统一落在 ``ARTIFACTS_DIR``，路径回传给模型并最终展示在聊天界面。
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
from pathlib import Path

from typing import Any

from paper_agent.core.constants import (
    ARTIFACTS_DIR,
    ATTACHMENTS_DIR,
    IMAGE_FILE_EXTENSIONS,
)
from paper_agent.services import session_artifacts
from paper_agent.services.background_assets import acquire_background, is_auto, is_cancelled
from paper_agent.services.paper_outline import format_workspace_file
from paper_agent.services.skills.base import Tool, ToolParam, ToolRegistry, ToolResult
from paper_agent.services.skills.code_sandbox import ExecutePythonTool
from paper_agent.utils.files import human_readable_size

# ---------------------------------------------------------------- Markdown 解析
# (kind, text)；kind: h1/h2/h3/p/li/quote/code/table/image
Block = tuple[str, str]
# 独占一行的图片：![说明](路径) —— 行内的图片仍按普通文本处理
_IMAGE_LINE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
# 图表题注：图 3-1 / 表 3-1 / Figure 3-1 / Table 3-1
_CAPTION_LINE = re.compile(r"^(图|表|figure|table)\s*\d+(?:[-–—.]\d+)?", re.IGNORECASE)
# 表题专用：论文规范是「表题在表上方」，需要单独认出来
_TABLE_CAPTION = re.compile(r"^(表格?|table)\s*\d+(?:[-–—.]\d+)?", re.IGNORECASE)


def parse_markdown_blocks(markdown: str) -> list[Block]:
    """把 Markdown 解析为线性块（够用即可，覆盖标题/列表/引用/代码/表格/段落）。"""
    blocks: list[Block] = []
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    index = 0
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(("p", " ".join(paragraph).strip()))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            index += 1
            buffer: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                buffer.append(lines[index])
                index += 1
            index += 1
            blocks.append(("code", "\n".join(buffer)))
            continue

        if not stripped:
            flush()
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            blocks.append((f"h{min(level, 3)}", heading.group(2).strip()))
            index += 1
            continue

        image = _IMAGE_LINE.match(stripped)
        if image:
            flush()
            blocks.append(("image", f"{image.group(1).strip()}\n{_clean_image_path(image.group(2))}"))
            index += 1
            continue

        list_item = re.match(r"^([-*+]|\d+[.)])\s+", stripped)
        if list_item:
            flush()
            # 缩进决定层级（li1/li2/li3）：PPT 要点分级、Word 缩进都要用
            indent = len(line) - len(line.lstrip(" \t"))
            level = min(indent // 2 + 1, 3)
            blocks.append((f"li{level}", _strip_inline(stripped[list_item.end():])))
            index += 1
            continue

        if stripped.startswith(">"):
            flush()
            blocks.append(("quote", _strip_inline(stripped.lstrip("> ").strip())))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|?[\s:\-|]+\|", lines[index + 1].strip()
        ):
            flush()
            table: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table.append(lines[index].strip())
                index += 1
            blocks.append(("table", "\n".join(table)))
            continue

        paragraph.append(_strip_inline(stripped))
        index += 1

    flush()
    return blocks


def _clean_image_path(raw: str) -> str:
    """去掉 Markdown 图片路径里的可选标题（``path "title"``）与尖括号包裹。"""
    path = re.sub(r'\s+"[^"]*"\s*$', "", raw.strip())
    path = re.sub(r"\s+'[^']*'\s*$", "", path.strip())
    return path.strip().strip("<>").strip()


def _split_image(text: str) -> tuple[str, str]:
    """把 image 块还原成 ``(说明文字, 图片路径)``。"""
    alt, _, path = (text or "").partition("\n")
    return alt.strip(), path.strip()


def _resolve_image(source: str) -> Path | None:
    """定位图片：先按原路径，再按文件名在产物（**本会话目录优先**）/ 附件里找。"""
    if not source:
        return None
    candidate = Path(source)
    if candidate.is_file():
        return candidate
    name = candidate.name
    if not name:
        return None
    for directory in (
        session_artifacts.dir_for(ARTIFACTS_DIR, create=False),
        ARTIFACTS_DIR,
        ATTACHMENTS_DIR,
    ):
        found = Path(directory) / name
        if found.is_file():
            return found
    return None


def _strip_inline(text: str) -> str:
    """去掉行内的 Markdown 标记（加粗 / 斜体 / 行内代码 / 链接）。"""
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    text = re.sub(r"(\*|_)(.*?)\1", r"\2", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text.strip()


def _needs_new_page(text: str, level: int, first_heading: bool, continuing: bool) -> bool:
    """这个标题要不要另起一页。

    论文规范：目录、每一章、参考文献、致谢、附录、声明都各自起新页。
    ``continuing``（续写模式）下不再对普通一级标题自动分页——追加的可能只是
    某一章的后半部分，但「目录 / 参考文献 / 致谢」这类标题仍然分页。
    """
    stripped = (text or "").strip()
    if stripped in PAGE_BREAK_TITLES:
        return True
    if _CHAPTER_PATTERN.search(stripped):
        return True
    return (not continuing) and level == 1 and not first_heading


def _split_citations(text: str) -> list[tuple[str, bool]]:
    """按参考文献角标切分正文，返回 ``[(片段, 是否要上标), ...]``。

    论文里引用要写成右上角标（如 ``...已有研究[1]表明``）。模型给的形式很杂：
    ``[1]``、``[1,3]``、``[2-5]``，也有直接写 Unicode 上标数字的 ``[¹⁰]``。
    这里统一还原成 ASCII 数字，整段（含方括号）渲染成真正的上标。
    """
    parts: list[tuple[str, bool]] = []
    position = 0
    for match in _CITATION.finditer(text or ""):
        if match.start() > position:
            parts.append((text[position:match.start()], False))
        body = match.group("body").translate(_SUPERSCRIPT_TO_ASCII)
        parts.append((f"[{body}]", True))
        position = match.end()
    if position < len(text or ""):
        parts.append((text[position:], False))
    return parts or [(text or "", False)]


def _add_body_paragraph(doc, text: str, style: str = "", superscript: bool = True):
    """写一段正文：参考文献角标渲染为右上角标。

    参考文献列表里的 ``[1]`` 是编号不是引用，``superscript=False`` 时不做上标。
    """
    paragraph = doc.add_paragraph(style=style or None)
    for fragment, is_citation in _split_citations(text):
        run = paragraph.add_run(fragment)
        if is_citation and superscript:
            run.font.superscript = True     # 右上角标
    return paragraph


def _doc_head_text(path: Path, limit: int = 30) -> str:
    """读已有 Word 文档开头几段（用于续写时判断排版风格）。失败返回空字符串。"""
    try:
        from docx import Document

        return "\n".join(p.text for p in Document(str(path)).paragraphs[:limit])
    except Exception:      # noqa: BLE001 - 读不了就按当次内容自行判断
        return ""


def _content_name(content: str) -> str:
    """从正文里派生的可读文件名（首个标题 / 首个非空行）。

    模型偶尔会漏传 ``filename``；直接用 ``output-时间戳`` 命名的话，
    用户在工作区里根本认不出哪个是哪个。
    """
    for line in (content or "").splitlines():
        text = _strip_inline(re.sub(r"^#{1,6}\s*", "", line.strip()))
        text = re.sub(r'[\\/:*?"<>|]', "_", text).strip(" ._")
        if text:
            return text[:40]
    return ""


def _safe_name(filename: str, default_ext: str, fallback: str = "") -> str:
    """清洗文件名并保证扩展名正确；``fallback`` 用于文件名缺失时兜底。"""
    name = re.sub(r'[\\/:*?"<>|]', "_", Path(filename or "").name).strip(" ._")
    if not name:
        name = re.sub(r'[\\/:*?"<>|]', "_", fallback or "").strip(" ._")
    if not name:
        name = f"output-{int(time.time())}"
    if not Path(name).suffix:
        name += default_ext
    return name


def normalize_path(path: str | Path) -> str:
    """路径归一化，用于比较（Windows 下大小写、正反斜杠都不敏感）。"""
    try:
        return str(Path(path).resolve()).lower()
    except OSError:
        return str(path).lower().replace("\\", "/")


def _unique_path(name: str, session_paths: set[str] | None = None) -> Path:
    """决定产物落盘路径（落在**本会话**的产物目录里）。

    同名文件若属于**本会话**（用户正在改的同一份文档），原地覆盖；
    否则退化为「-1、-2」后缀，避免堆一堆副本。
    """
    return session_artifacts.unique_path(
        name, base=ARTIFACTS_DIR, taken=session_paths or ()
    )


# ---------------------------------------------------------------- Word
# 论文排版参数：正文宋体小四 / 标题黑体（三号·四号·小四）/ 表格五号
THESIS_BODY_FONT = ("Times New Roman", "宋体", 12)      # (西文, 中文, 磅值)
THESIS_HEADING_FONTS = {1: ("黑体", 16), 2: ("黑体", 14), 3: ("黑体", 12)}
THESIS_CODE_FONT = ("Consolas", "Consolas", 10.5)
THESIS_TABLE_FONT = ("Times New Roman", "宋体", 10.5)
THESIS_FIRST_LINE_INDENT = 24          # 小四两字符 = 24 磅

# 「目录」标题的写法：命中就换成 Word 真目录域
TOC_HEADINGS = {"目录", "目次", "contents", "table of contents"}

# 这些标题必须另起一页（论文规范：目录、各章、参考文献、致谢等各自分页）
PAGE_BREAK_TITLES = {
    "目录", "目次", "参考文献", "致谢", "附录", "声明",
    "摘要", "ABSTRACT", "Abstract", "中文摘要", "英文摘要",
    "原创性声明", "学位论文原创性声明", "学位论文版权使用授权书",
    "攻读学位期间发表的论文与研究成果",
}
# 正文里的参考文献角标。模型经常直接写 Unicode 上标数字（`[¹]`、`[¹⁰]`），
# 那样只有数字是小的、方括号跟正文一样大，必须统一转成整段真正的上标。
_SUPERSCRIPT_DIGITS = "⁰¹²³⁴⁵⁶⁷⁸⁹"
_SUPERSCRIPT_TO_ASCII = str.maketrans(_SUPERSCRIPT_DIGITS, "0123456789")
_CITATION = re.compile(
    r"[\[【［]\s*"
    r"(?P<body>[0-9" + _SUPERSCRIPT_DIGITS + r"]+"
    r"(?:\s*[,\-–、~～]\s*[0-9" + _SUPERSCRIPT_DIGITS + r"]+)*)"
    r"\s*[\]】］]"
)

# 判断「论文」的特征词
THESIS_MARKERS = (
    "摘要", "关键词", "参考文献", "绪论", "文献综述", "致谢",
    "开题", "毕业论文", "学位论文", "目录", "研究背景",
)
_CHAPTER_PATTERN = re.compile(r"第\s*[一二三四五六七八九十百零〇\d]+\s*[章节篇]")


def detect_doc_style(markdown: str, filename: str = "") -> str:
    """按内容判断是「论文」还是普通文档，返回 ``"thesis"`` / ``"plain"``。

    命中任一论文特征即按论文排版：出现「第X章」、或同时出现「摘要 + 关键词」、
    或文件名与正文里累计出现 3 个以上论文特征词。
    """
    text = markdown or ""
    if _CHAPTER_PATTERN.search(text):
        return "thesis"
    if "摘要" in text and "关键词" in text:
        return "thesis"
    scope = f"{filename}\n{text[:3000]}"
    hits = sum(1 for marker in THESIS_MARKERS if marker in scope)
    return "thesis" if hits >= 3 else "plain"


def resolve_doc_style(style: str, markdown: str, filename: str = "") -> str:
    """把工具传入的 ``doc_type`` 归一化成 ``"thesis"`` / ``"plain"``。"""
    normalised = (style or "auto").strip().lower()
    if normalised in {"thesis", "论文", "paper"}:
        return "thesis"
    if normalised in {"plain", "普通", "normal", "doc"}:
        return "plain"
    return detect_doc_style(markdown, filename)


def _apply_font(target, ascii_name: str, east_asian: str, size: float, bold: bool = False) -> None:
    """设置字体（含中文 eastAsia），并清掉主题字体属性，保证 Word 里真正生效。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    target.font.name = ascii_name
    target.font.size = Pt(size)
    target.font.bold = bold

    element = getattr(target, "element", None)
    if element is None:                       # Run 用私有属性（style 与 run 通用）
        element = target._element             # noqa: SLF001
    fonts = element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), east_asian)


def _init_docx_styles(doc, thesis: bool) -> None:
    """论文：正文宋体小四 + 1.5 倍行距 + 首行缩进 2 字符，标题黑体分级。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    normal = doc.styles["Normal"]
    if not thesis:
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(11)
        return

    _apply_font(normal, *THESIS_BODY_FONT)
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(THESIS_FIRST_LINE_INDENT)

    for level, (font_name, size) in THESIS_HEADING_FONTS.items():
        style = doc.styles[f"Heading {level}"]
        _apply_font(style, "Times New Roman", font_name, size, bold=True)
        style.font.color.rgb = RGBColor(0, 0, 0)       # 去掉 Word 默认的蓝色标题
        heading_format = style.paragraph_format
        heading_format.line_spacing = 1.5
        heading_format.first_line_indent = Pt(0)
        heading_format.space_before = Pt(18 if level == 1 else 12)
        heading_format.space_after = Pt(12 if level == 1 else 6)
        if level == 1:
            heading_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_toc_heading(doc, text: str):
    """「目录」标题 + Word 目录域，返回标题段落。

    标题故意不用 Heading 样式，避免目录把自己也列进去；真正的条目由 Word 生成。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    paragraph = doc.add_paragraph()
    format_ = paragraph.paragraph_format
    format_.first_line_indent = Pt(0)
    format_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_.space_before = Pt(18)
    format_.space_after = Pt(12)
    run = paragraph.add_run(text)
    _apply_font(run, "Times New Roman", "黑体", THESIS_HEADING_FONTS[1][1], bold=True)
    _add_toc_field(doc)
    return paragraph


def _three_line_borders(table) -> None:
    """把表格做成论文要求的三线表：顶线 + 表头下线 + 底线，无竖线、无内部横线。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tblBorders")
    for edge, size, style in (
        ("top", "12", "single"),          # 1.5 磅
        ("bottom", "12", "single"),
        ("left", "0", "none"),
        ("right", "0", "none"),
        ("insideH", "0", "none"),
        ("insideV", "0", "none"),
    ):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), style)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), "000000")
        borders.append(element)
    table._tbl.tblPr.append(borders)              # noqa: SLF001

    for cell in table.rows[0].cells:              # 表头行下线（0.75 磅）
        cell_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        cell_borders.append(bottom)
        cell._tc.get_or_add_tcPr().append(cell_borders)   # noqa: SLF001

    header_row = table.rows[0]._tr                # noqa: SLF001 - 跨页时重复表头
    properties = header_row.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _style_caption(doc, text: str, thesis: bool) -> None:
    """图表题注：居中、五号、加粗、不缩进、与图表贴紧。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    paragraph = doc.add_paragraph()
    format_ = paragraph.paragraph_format
    format_.first_line_indent = Pt(0)
    format_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_.space_before = Pt(3)
    format_.space_after = Pt(3)
    run = paragraph.add_run(text)
    if thesis:
        _apply_font(run, "Times New Roman", "宋体", THESIS_TABLE_FONT[2], bold=True)
    else:
        run.bold = True


def _add_image(doc, source: str, caption: str, thesis: bool) -> None:
    """插入配图（等比缩放到正文宽度内）并在图下加题注。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, Pt

    resolved = _resolve_image(source)
    if resolved is None:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"［图片未找到：{source}］")
        run.font.size = Pt(10)
        if thesis:
            _apply_font(run, "Times New Roman", "宋体", 10.5)
        return

    try:
        shape = doc.add_picture(str(resolved))
    except Exception:      # noqa: BLE001 - 图片损坏 / 格式不支持时给占位文字
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.add_run(f"［图片无法插入：{resolved.name}］")
        return

    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    limit = Emu(int(usable * 0.92))           # 留一点边距，避免顶到页边
    if shape.width > limit:
        ratio = limit / shape.width
        shape.width = int(limit)
        shape.height = int(shape.height * ratio)

    picture_paragraph = doc.paragraphs[-1]     # add_picture 建在最后一个段落
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = Pt(0)

    if caption:
        _style_caption(doc, caption, thesis)


def _add_toc_field(doc) -> None:
    """插入 Word 目录域：在 Word 里按 F9（或打开时自动）即可生成真实页码目录。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = 0
    for tag, attrib, content in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, ' TOC \\o "1-3" \\h \\z \\u '),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "（在 Word 中按 F9 或右键「更新域」生成目录）"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attrib.items():
            element.set(qn(key), value)
        if content is not None:
            element.text = content
        run = paragraph.add_run()
        run._r.append(element)            # noqa: SLF001


def _enable_update_fields(doc) -> None:
    """让 Word 打开文档时更新域（否则目录要手动 F9）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    doc.settings.element.append(element)


def _add_page_number(doc) -> None:
    """页脚居中页码（论文正文一般要求有页码）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, attrib, content in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "1"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attrib.items():
            element.set(qn(key), value)
        if content is not None:
            element.text = content
        run = paragraph.add_run()
        run._r.append(element)            # noqa: SLF001
        _apply_font(run, "Times New Roman", "宋体", 10.5)


def _is_manual_toc_line(text: str) -> bool:
    """模型手写的目录行（``第一章 绪论\t6`` / ``1.1 背景……6``）——不要写进正文。"""
    stripped = (text or "").strip()
    if not stripped:
        return False
    return bool(re.search(r"(?:[\t.·…]{1,}|\s{2,})\s*\d{1,4}\s*$", stripped))


def _add_table(doc, rows: list[list[str]], thesis: bool, caption: str = "") -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    if caption:
        _style_caption(doc, caption, thesis)

    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    if thesis:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _three_line_borders(table)

    for r, row in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            if not thesis:
                continue
            for paragraph in cell.paragraphs:
                cell_format = paragraph.paragraph_format
                cell_format.first_line_indent = Pt(0)
                cell_format.line_spacing = 1.0
                cell_format.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    _apply_font(run, *THESIS_TABLE_FONT, bold=(r == 0))

    if thesis:                                          # 避免表格与后续段落粘连
        doc.add_paragraph()


def build_docx(
    markdown: str, path: Path, style: str = "auto", append: bool = False
) -> None:
    """把 Markdown 写成 Word 文档。

    Args:
        style: ``thesis`` 按学位论文排版（宋体正文 / 黑体标题 / 1.5 倍行距 /
            首行缩进 2 字符 / 三线表 / 图表题注 / 目录域 / 页脚页码）；
            ``plain`` 普通文档；``auto`` 按内容自动判断。
        append: 往已存在的文档尾部续写（长文分段写入时用），``False`` 则新建 / 覆盖。
    """
    from docx import Document
    from docx.shared import Pt

    thesis = resolve_doc_style(style, markdown, path.name) == "thesis"

    continuing = append and path.exists()
    if continuing:
        doc = Document(str(path))     # 接着已有文档写，保留之前的内容与页码
    else:
        doc = Document()
    _init_docx_styles(doc, thesis)
    if thesis and not continuing:
        _enable_update_fields(doc)
        _add_page_number(doc)

    blocks = parse_markdown_blocks(markdown)
    index = 0
    seen_heading = False
    in_references = False        # 处于「参考文献」章节：这里的 [1] 是编号，不做角标
    while index < len(blocks):
        kind, text = blocks[index]
        following = blocks[index + 1] if index + 1 < len(blocks) else ("", "")

        if kind in ("h1", "h2", "h3"):
            # 表题被写成标题行时不单独渲染，交给表格挂在表上方
            if _TABLE_CAPTION.match(text.strip()) and following[0] == "table":
                index += 1
                continue
            level = int(kind[1])
            new_page = thesis and _needs_new_page(text, level, not seen_heading, continuing)
            seen_heading = True
            in_references = text.strip() in {"参考文献", "references", "REFERENCES"}
            if thesis and text.strip() in TOC_HEADINGS:
                heading = _add_toc_heading(doc, text.strip())
            else:
                heading = doc.add_heading(text, level=level)
            if new_page:
                # 用「段前分页」属性而不是插入分页符段落：后者在 Word 里会显示成
                # 一整行「…………分页符…………」，排版上也不干净
                heading.paragraph_format.page_break_before = True
        elif kind.startswith("li"):
            paragraph = _add_body_paragraph(
                doc, text, style="List Bullet", superscript=not in_references
            )
            if thesis:
                # 二级及以下要点额外缩进，避免和一级混在一起
                level = int(kind[2:])
                if level > 1:
                    paragraph.paragraph_format.left_indent = Pt(18 * (level - 1))
        elif kind == "quote":
            _add_body_paragraph(
                doc, text, style="Intense Quote", superscript=not in_references
            )
        elif kind == "code":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(text)
            if thesis:
                paragraph.paragraph_format.first_line_indent = Pt(0)
                _apply_font(run, *THESIS_CODE_FONT)
            else:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif kind == "image":
            alt, source = _split_image(text)
            caption = alt if _CAPTION_LINE.match(alt) else ""
            if following[0] == "p" and _CAPTION_LINE.match(following[1].strip()):
                caption = following[1].strip()         # 图题在图下方
                index += 1
            _add_image(doc, source, caption, thesis)
        elif kind == "table":
            rows = _table_rows(text)
            if rows:
                # 表题在表上方：上一轮的「表 x-y」行会作为题注一起输出
                caption = _pending_caption(blocks, index)
                if not caption and following[0] == "p" and _TABLE_CAPTION.match(
                    following[1].strip()
                ):
                    # 模型偶尔把表题写在表下面：挪到表上方（论文规范）
                    caption = following[1].strip()
                    index += 1
                _add_table(doc, rows, thesis, caption)
        elif kind == "p" and _is_manual_toc_line(text) and _in_toc_section(blocks, index):
            pass                                       # 手写目录行交给 Word 目录域
        elif kind == "p" and _TABLE_CAPTION.match(text.strip()) and following[0] == "table":
            pass                                       # 表题在下一轮随表格一起输出
        else:
            _add_body_paragraph(
                doc, text, superscript=not in_references   # 正文里的引用是右上角标
            )
        index += 1

    doc.save(str(path))


def _pending_caption(blocks: list[Block], index: int) -> str:
    """表格上方的「表 x-y 标题」（若有）作为题注。

    模型有时把表题写成标题行（`#### 表 3-3 …`），有时和表格之间还夹一段空话，
    因此标题行也算、并往上多看一两个块。
    """
    for position in range(index - 1, max(-1, index - 3), -1):
        previous_kind, previous_text = blocks[position]
        candidate = previous_text.strip()
        if previous_kind in ("p", "h1", "h2", "h3") and _TABLE_CAPTION.match(candidate):
            return candidate
    return ""


def _in_toc_section(blocks: list[Block], index: int) -> bool:
    """往上找最近的一级标题，判断当前是否处于「目录」章节里。"""
    for position in range(index - 1, -1, -1):
        kind, text = blocks[position]
        if kind == "h1":
            return text.strip() in TOC_HEADINGS
    return False


def _table_rows(table_text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in table_text.split("\n"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if re.match(r"^[\s:\-]+$", "".join(cells)):
            continue
        rows.append(cells)
    return rows


# ---------------------------------------------------------------- PPT
def build_pptx(
    markdown: str,
    path: Path,
    background: str = "",
    backgrounds: list[str] | None = None,
) -> None:
    """把 Markdown 大纲生成 16:9 演示文稿。

    版式：封面页（渐变底）→ 目录页（章节 ≥5 时自动生成）→ 内容页
    （标题 + 主题色装饰条 + 分级要点 + 页脚页码）。内容过多会自动分页。

    Args:
        markdown: ``#`` 为一页标题，``##`` 为页内小节，``-`` / 缩进为分级要点；
            章节下单独一行 ``[背景图: 文件名或描述]`` 可给该节单独换背景
        path: 输出文件路径
        background: 单张背景图（路径）；兼容旧用法
        backgrounds: 多张背景图路径，按章节顺序轮流使用，避免每页一样

    背景处理方式会按页轮换（整页蒙版 / 左侧渐变 / 右侧竖条）并变换裁切焦点，
    因此只有一张图时也不会每页长得完全一样。
    """
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
    prs.slide_height = Inches(SLIDE_HEIGHT_INCHES)
    blank = prs.slide_layouts[6]

    pool = list(backgrounds or [])
    if background and background not in pool:
        pool.insert(0, background)
    backdrops = [item for item in (_resolve_image(value) for value in pool) if item]

    sections, preamble = _pptx_sections(parse_markdown_blocks(markdown))
    # 「目录」由系统生成：手写的目录章节既不占页也不进目录
    sections = [item for item in sections if not _is_navigation_title(item["title"])]
    cover_title, subtitle, sections = _pptx_cover_info(sections, preamble, path)

    # 目录只列一级标题，且按「原始章节」计数：自动分页出来的（续）页不再各算一条
    pages = _pptx_split_long_sections(sections)
    toc_entries = _pptx_toc_entries(pages)
    backdrop = backdrops[0] if backdrops else None

    _pptx_cover(prs.slides.add_slide(blank), cover_title, subtitle, backdrop)
    if len(toc_entries) >= TOC_MIN_SECTIONS:
        _pptx_toc(prs.slides.add_slide(blank), toc_entries, cover_title, backdrop)
    for index, section in enumerate(pages, start=1):
        image = None
        # 本节单独指定的背景优先；否则按章节序号在背景池里轮换
        marker = section.get("background") or ""
        if marker:
            image = _resolve_image(marker)
        if image is None and backdrops:
            image = backdrops[(section.get("origin", index - 1)) % len(backdrops)]
        style, focus = _background_variant(
            index, bool(section.get("image") or section.get("table"))
        )
        _pptx_section(
            prs.slides.add_slide(blank), section, index, len(pages), cover_title,
            image, style=style, focus=focus,
        )

    if not sections and not cover_title:      # 完全空内容时也给一页提示
        slide = prs.slides.add_slide(blank)
        _pptx_page_background(slide, backdrop, cover=False)
        _pptx_text_box(
            slide, TEXT_LEFT, Inches(3), Inches(10), Inches(1),
            "（内容为空）", size=20, color=PPTX_PALETTE["muted"],
        )

    prs.save(str(path))


# ---------------------------------------------------------------- PPT 版式
SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5
TEXT_LEFT = 0.95
TEXT_WIDTH = 11.4
BODY_TOP = 1.85
BODY_HEIGHT = 4.55
MAX_ITEMS_PER_SLIDE = 7          # 超过就自动分页
TOC_MIN_SECTIONS = 5             # 章节够多才自动加目录页
MAX_BACKDROPS = 4                # 一份 PPT 最多「联网生成」几张背景（每张都要出图，成本高）
MAX_LOCAL_BACKDROPS = 8          # 本地已有图片的上限（不花出图额度，可以多放几张）
# 章节级背景指令：`[背景图: 本地文件名.jpg]` 或 `[背景图: 画面描述]`
BACKGROUND_MARKER = re.compile(r"^\[背景图?\s*[:：]\s*(.+?)\]$")
# 背景处理方式的轮换：整页蒙版 → 左侧渐变 → 右侧竖条 → 左侧渐变 → 整页蒙版
BACKGROUND_CYCLE = ("full", "scrim", "side", "scrim", "full")
BACKGROUND_FOCUS = ("center", "left", "right")
PPTX_TABLE_MAX_ROWS = 9          # 单页表格最多行数（超出会在页面上注明）
PPTX_TABLE_MAX_COLS = 6          # 单页表格最多列数
TOC_MAX_ENTRIES = 12             # 目录页最多列几条（再多会溢出页面）

PPTX_PALETTE = {
    "accent": "1F4E79",
    "accent_dark": "12324F",
    "accent_soft": "E7EFF8",
    "text": "1F2937",
    "muted": "7A8798",
    "cover_text": "FFFFFF",
    "page_bg": "FAFBFD",
}
PPTX_FONT = "微软雅黑"
PPTX_BULLETS = ("●", "○", "▪")
PPTX_ITEM_SIZES = {"bullet": 20, "sub": 22, "text": 18, "quote": 16, "code": 13}


# 导航类标题：目录是系统自动生成的，不该自己占一页，也不该出现在目录里
PPTX_NAVIGATION_TITLES = {"目录", "目次", "contents", "table of contents"}


def _norm_title(text: str) -> str:
    """标题归一化（去空白与标点、转小写），用于相似度比较。"""
    return re.sub(r"[\s\-_·:：、,，.。()（）\[\]【】/\\]+", "", (text or "").strip().lower())


def _similar_title(left: str, right: str, threshold: float = 0.6) -> bool:
    from difflib import SequenceMatcher

    first, second = _norm_title(left), _norm_title(right)
    if not first or not second:
        return False
    return SequenceMatcher(None, first, second).ratio() >= threshold


def _title_from_filename(path: Path) -> str:
    """从文件名推测讲题，去掉 `_v3`、`-2` 这类版本后缀。"""
    stem = re.sub(r"[-_\s]*(?:v|V)\d+$", "", path.stem)
    return re.sub(r"[-_\s]*\d+$", "", stem) or path.stem


def _is_navigation_title(title: str) -> bool:
    return (title or "").strip().lower() in PPTX_NAVIGATION_TITLES


def _deck_topic(markdown: str, stem: str = "") -> str:
    """取讲题：第一个一级标题，没有就用文件名。"""
    for kind, text in parse_markdown_blocks(markdown or ""):
        if kind == "h1" and text.strip():
            return text.strip()
    return (stem or "").strip() or "学术汇报"


def _pptx_sections(blocks: list[Block]) -> tuple[list[dict], list[str]]:
    """把线性块切成「每页一节」，并单独收集正文前的零散段落（做封面副标题）。"""
    sections: list[dict] = []
    preamble: list[str] = []
    current: dict | None = None

    for kind, text in blocks:
        if kind == "h1":
            current = {"title": text.strip(), "items": [], "image": "", "table": []}
            sections.append(current)
            continue
        if current is None:
            if kind == "p":
                preamble.append(text.strip())
            continue
        if kind == "p":
            # `[背景图: xxx]` 是给这一节换背景的指令，不是正文
            marker = BACKGROUND_MARKER.match(text.strip())
            if marker:
                current["background"] = marker.group(1).strip()
                continue
            current["items"].append(("text", 0, text.strip()))
        elif kind == "h2":
            current["items"].append(("sub", 0, text.strip()))
        elif kind.startswith("li"):
            current["items"].append(("bullet", int(kind[2:]) - 1, text.strip()))
        elif kind == "image":
            _alt, source = _split_image(text)
            current["image"] = source
        elif kind == "table":
            current["table"] = _table_rows(text)
        elif kind == "code":
            current["items"].append(("code", 0, text.strip()))
        elif kind == "quote":
            current["items"].append(("quote", 0, text.strip()))
        else:
            current["items"].append(("text", 0, text.strip()))
    return sections, preamble


def _pptx_cover_info(
    sections: list[dict], preamble: list[str], path: Path
) -> tuple[str, str, list[dict]]:
    """确定封面标题与副标题。

    **正文内容绝不丢弃**：
    - 只有「标题 + 一句短话」的第一节才整节当作封面；
    - 带要点 / 长段落 / 图表的第一节一律保留为正文页，即使它的标题和讲题很像
      （那种情况只把标题借给封面用，并标记不进目录，避免目录出现重复条目）。
    """
    cover_title, subtitle = "", ""
    if sections:
        first = sections[0]
        short_texts = [
            text for kind, _level, text in first["items"] if kind == "text" and text.strip()
        ]
        has_points = (
            any(kind in {"bullet", "sub"} for kind, _level, _text in first["items"])
            or bool(first["image"])
            or bool(first["table"])
        )
        only_short = len(short_texts) <= 1 and all(len(text) <= 60 for text in short_texts)
        if not has_points and only_short:
            cover_title = first["title"]
            subtitle = short_texts[0] if short_texts else ""
            sections = sections[1:]
        elif _similar_title(first["title"], _title_from_filename(path)):
            # 这一节的标题其实就是讲题（模型常把它又写成一个章节）：
            # 拿它当封面标题更好看，但内容留着，只是不再进目录
            cover_title = first["title"]
            subtitle = next((text for text in short_texts if len(text) <= 60), "")
            first["no_toc"] = True
    if not cover_title:
        cover_title = _title_from_filename(path) or "演示文稿"
    if not subtitle:
        subtitle = next((text for text in preamble if len(text) <= 60), "")
    return cover_title, subtitle, sections


def _pptx_split_long_sections(sections: list[dict]) -> list[dict]:
    """一页放不下的要点自动拆成续页。

    ``origin`` 记录该页来自哪一节，供目录算「首屏页码」（续页不重复计入）。
    """
    result: list[dict] = []
    for origin, section in enumerate(sections):
        items = section["items"]
        if len(items) <= MAX_ITEMS_PER_SLIDE:
            section["origin"] = origin
            result.append(section)
            continue
        chunks = [
            items[start : start + MAX_ITEMS_PER_SLIDE]
            for start in range(0, len(items), MAX_ITEMS_PER_SLIDE)
        ]
        for index, chunk in enumerate(chunks):
            copy = dict(section)
            copy["origin"] = origin
            copy["title"] = section["title"] if index == 0 else f"{section['title']}（续）"
            copy["items"] = chunk
            if index:                       # 图表只跟第一页
                copy["image"] = ""
                copy["table"] = []
            result.append(copy)
    return result


def _pptx_toc_entries(pages: list[dict]) -> list[tuple[str, int]]:
    """目录条目：``(标题, 首屏页码)``。

    - 只列一级标题，自动分页出来的「（续）」页不重复计入；
    - 页码取该节第一页的页序号，与页脚页码一致；
    - 标了 ``no_toc`` 的节（标题已被用作封面）不进目录。
    """
    entries: list[tuple[str, int]] = []
    seen: set[int] = set()
    for page_no, page in enumerate(pages, start=1):
        origin = page.get("origin", page_no)
        if origin in seen:
            continue
        seen.add(origin)
        if page.get("no_toc"):
            continue
        entries.append((page["title"], page_no))
    return entries


# ---------------------------------------------------------------- 绘图辅助
def _pptx_new_element(parent, tag: str, **attrs):
    """在指定节点下追加一个 XML 子节点（lxml 自带 makeelement，无需额外依赖）。"""
    from pptx.oxml.ns import qn

    element = parent.makeelement(qn(tag), {key: str(value) for key, value in attrs.items()})
    parent.append(element)
    return element


def _pptx_shape(slide, shape_type, left, top, width, height, color: str | None = None):
    """加一个无边框、无阴影的形状（PPT 单位可直接用 Inches(...)）。"""
    from pptx.dml.color import RGBColor

    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.shadow.inherit = False
    shape.line.fill.background()
    if color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
    return shape


def _pptx_alpha(shape, percent: int) -> None:
    """给纯色填充加透明度（python-pptx 没有 API，改 XML）。"""
    from pptx.oxml.ns import qn

    solid = shape._element.spPr.find(qn("a:solidFill"))      # noqa: SLF001
    if solid is None:
        return
    color = solid.find(qn("a:srgbClr"))
    if color is None:
        return
    _pptx_new_element(color, "a:alpha", val=int(percent * 1000))


def _pptx_gradient(shape, start: str, end: str, angle: float = 45.0) -> None:
    from pptx.dml.color import RGBColor

    shape.fill.gradient()
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = RGBColor.from_string(start)
    stops[0].position = 0.0
    stops[1].color.rgb = RGBColor.from_string(end)
    stops[1].position = 1.0
    shape.fill.gradient_angle = angle


def _pptx_run_font(run, name: str, size: float, *, bold: bool = False, color: str = "") -> None:
    """设置字体；中文必须写 ``a:ea``，否则 PowerPoint 里回退成默认宋体。"""
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    r_pr = run._r.get_or_add_rPr()                            # noqa: SLF001
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        latin = r_pr.find(qn("a:latin"))
        ea = r_pr.makeelement(qn("a:ea"), {})
        if latin is not None:
            latin.addnext(ea)                                # 保持标签顺序合法
        else:
            r_pr.append(ea)
    ea.set("typeface", name)


def _pptx_bullet(paragraph, char: str, color: str, level: int) -> None:
    """显式设置项目符号与悬挂缩进（不依赖母版的列表样式）。

    子元素顺序要符合 OOXML 规定：buClr → buSzPct → buFont → buChar。
    """
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    properties = paragraph._p.get_or_add_pPr()                # noqa: SLF001
    properties.set("marL", str(int(Pt(20 + level * 24))))
    properties.set("indent", str(int(-Pt(20))))
    color_element = _pptx_new_element(properties, "a:buClr")
    _pptx_new_element(color_element, "a:srgbClr", val=color)
    _pptx_new_element(properties, "a:buSzPct", val=90000)     # 符号取正文 90%
    _pptx_new_element(properties, "a:buFont", typeface="Arial")
    _pptx_new_element(properties, "a:buChar", char=char)


def _pptx_no_bullet(paragraph) -> None:
    from pptx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()                # noqa: SLF001
    _pptx_new_element(properties, "a:buNone")


def _pptx_text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: float,
    color: str,
    bold: bool = False,
    font: str = PPTX_FONT,
    align=None,
):
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _pptx_run_font(run, font, size, bold=bold, color=color)
    return box


def _background_variant(index: int, has_visual: bool) -> tuple[str, str]:
    """按页序轮换背景的处理方式与裁切焦点。

    只有一张背景图时，靠「整页蒙版 / 左侧渐变 / 右侧竖条」+ 三种裁切焦点
    产生变化，避免整份 PPT 每页长得一模一样；有图表的那页右侧要放内容，
    所以不用右侧竖条。
    """
    style = BACKGROUND_CYCLE[(index - 1) % len(BACKGROUND_CYCLE)]
    if style == "side" and has_visual:
        style = "scrim"
    focus = BACKGROUND_FOCUS[(index - 1) % len(BACKGROUND_FOCUS)]
    return style, focus


def _pptx_page_background(
    slide, image: Path | None, *, cover: bool, style: str = "full", focus: str = "center"
) -> None:
    """铺底：有背景图按指定样式铺，否则用主题色渐变 / 浅底 + 装饰条。

    style:
        ``full``  整页铺满 + 蒙版（默认，最稳）
        ``scrim`` 整页铺满 + 左侧白到透明的渐变蒙版（更有层次）
        ``side``  图片只占右侧竖条，左半页留白写正文
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    width = Inches(SLIDE_WIDTH_INCHES)
    height = Inches(SLIDE_HEIGHT_INCHES)

    if image is not None and style == "side" and not cover:
        _pptx_cover_picture(
            slide, image, Inches(8.15), 0, Inches(SLIDE_WIDTH_INCHES - 8.15), height,
            focus=focus,
        )
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["page_bg"])
        stripe = _pptx_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), height)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["accent"])
        return

    if image is not None:
        _pptx_cover_picture(slide, image, 0, 0, width, height, focus=focus)
        veil = _pptx_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, width, height)
        veil.fill.solid()
        veil.fill.fore_color.rgb = RGBColor.from_string(
            PPTX_PALETTE["accent_dark"] if cover else "FFFFFF"
        )
        _pptx_alpha(veil, 58 if cover else (72 if style == "scrim" else 78))
        if style == "scrim" and not cover:      # 左侧再压一层实色，保证正文可读
            scrim = _pptx_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, Inches(8.6), height)
            scrim.fill.solid()
            scrim.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            _pptx_alpha(scrim, 62)
        return

    if cover:
        base = _pptx_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, width, height)
        _pptx_gradient(base, PPTX_PALETTE["accent"], PPTX_PALETTE["accent_dark"], angle=35.0)
        for left, top, size, alpha in (       # 装饰圆只放在页内，避免编辑器里看起来残缺
            (9.35, 0.55, 3.4, 12),
            (11.0, 3.35, 2.3, 10),
            (8.6, 4.55, 1.5, 8),
        ):
            blob = _pptx_shape(
                slide, MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(size), Inches(size)
            )
            blob.fill.solid()
            blob.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            _pptx_alpha(blob, alpha)
        return

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["page_bg"])
    stripe = _pptx_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["accent"])
    corner = _pptx_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.15), Inches(5.62), Inches(3.05), Inches(1.85),
    )
    corner.fill.solid()
    corner.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["accent_soft"])
    corner.shadow.inherit = False


def _pptx_cover_picture(slide, image: Path, left, top, width, height, focus: str = "center") -> None:
    """按「覆盖」方式把图铺满指定矩形：等比放大后裁掉多余部分，不拉伸变形。

    ``focus`` 决定裁掉哪一侧，同一张图裁出不同区域，观感就不会每页都一样。
    """
    picture = slide.shapes.add_picture(str(image), left, top, width=width)
    natural = picture.width / max(picture.height, 1)      # 原图宽高比
    target = width / max(height, 1)                       # 目标区域宽高比
    margin = max(0.0, 1 - min(natural / target, target / natural)) / 2

    picture.crop_left = picture.crop_right = 0.0
    picture.crop_top = picture.crop_bottom = 0.0
    if natural > target:                                  # 图更宽：裁左右
        if focus == "left":                               # 保留左侧内容
            picture.crop_left, picture.crop_right = 0.0, margin * 2
        elif focus == "right":                            # 保留右侧内容
            picture.crop_left, picture.crop_right = margin * 2, 0.0
        else:
            picture.crop_left = picture.crop_right = margin
    elif natural < target:                                # 图更高：裁上下
        if focus == "left":
            picture.crop_top, picture.crop_bottom = 0.0, margin * 2
        elif focus == "right":
            picture.crop_top, picture.crop_bottom = margin * 2, 0.0
        else:
            picture.crop_top = picture.crop_bottom = margin

    picture.width = width
    picture.height = height
    picture.left = left
    picture.top = top


def _pptx_heading(slide, text: str) -> None:
    """页标题（左侧色块 + 标题 + 细分隔线）。"""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    _pptx_text_box(
        slide, TEXT_LEFT, 0.62, TEXT_WIDTH, 1.0, text,
        size=30, color=PPTX_PALETTE["text"], bold=True,
    )
    bar = _pptx_shape(
        slide, MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(0.78), Inches(0.16), Inches(0.62)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["accent"])
    line = _pptx_shape(
        slide, MSO_SHAPE.RECTANGLE, Inches(TEXT_LEFT), Inches(1.5), Inches(TEXT_WIDTH), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string(PPTX_PALETTE["accent_soft"])


def _pptx_footer(slide, deck_title: str, page: int, total: int) -> None:
    from pptx.enum.text import PP_ALIGN

    _pptx_text_box(slide, TEXT_LEFT, 6.72, 8.0, 0.4, deck_title, size=11, color=PPTX_PALETTE["muted"])
    if page:
        _pptx_text_box(
            slide, 11.4, 6.72, 1.2, 0.4, f"{page} / {total}",
            size=11, color=PPTX_PALETTE["muted"], align=PP_ALIGN.RIGHT,
        )


def _pptx_cover(slide, title: str, subtitle: str, image: Path | None) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    _pptx_page_background(slide, image, cover=True)
    text_color = PPTX_PALETTE["cover_text"]
    _pptx_text_box(
        slide, TEXT_LEFT, 2.55, TEXT_WIDTH - 1.0, 1.6, title,
        size=42, color=text_color, bold=True,
    )
    accent = _pptx_shape(
        slide, MSO_SHAPE.RECTANGLE, Inches(TEXT_LEFT), Inches(4.18), Inches(1.5), Inches(0.07)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string("7FB2E5")
    if subtitle:
        _pptx_text_box(
            slide, TEXT_LEFT, 4.45, TEXT_WIDTH - 2.0, 0.9, subtitle, size=18, color="D8E2EE"
        )
    _pptx_text_box(
        slide, TEXT_LEFT, 6.2, 6.0, 0.5, time.strftime("%Y 年 %m 月 %d 日"),
        size=13, color="B9C9DC",
    )


def _pptx_toc(
    slide, entries: list[tuple[str, int]], deck_title: str, image: Path | None
) -> None:
    """目录页：条目为 ``(标题, 首屏页码)``，页码与正文页脚保持一致。"""
    from pptx.util import Inches

    _pptx_page_background(slide, image, cover=False)
    _pptx_heading(slide, "目录")
    hidden = max(0, len(entries) - TOC_MAX_ENTRIES)
    entries = entries[:TOC_MAX_ENTRIES]           # 条目过多会溢出页面
    columns = 2 if len(entries) > 6 else 1
    rows = (len(entries) + columns - 1) // columns
    column_width = TEXT_WIDTH / columns
    for index, (title, _page_no) in enumerate(entries):
        column, row = divmod(index, rows)
        _pptx_text_box(
            slide,
            TEXT_LEFT + column * column_width,
            BODY_TOP + row * 0.62,
            column_width - 0.4,
            0.5,
            title,                      # 目录条目不显示页码前缀
            size=17,
            color=PPTX_PALETTE["accent"] if column == 0 else PPTX_PALETTE["text"],
        )
    if hidden:
        _pptx_text_box(
            slide,
            TEXT_LEFT,
            BODY_TOP + rows * 0.62 + 0.05,
            TEXT_WIDTH,
            0.4,
            f"（还有 {hidden} 个章节见正文）",
            size=13,
            color=PPTX_PALETTE["muted"],
        )
    _pptx_footer(slide, deck_title, 0, 0)


def _pptx_section(
    slide,
    section: dict,
    index: int,
    total: int,
    deck_title: str,
    image: Path | None,
    style: str = "full",
    focus: str = "center",
) -> None:
    from pptx.util import Inches

    _pptx_page_background(slide, image, cover=False, style=style, focus=focus)
    _pptx_heading(slide, section["title"])

    has_visual = bool(section["image"]) or bool(section["table"])
    text_width = 6.0 if has_visual else TEXT_WIDTH
    right_left = 7.25

    _pptx_body(slide, section["items"], TEXT_LEFT, BODY_TOP, text_width, BODY_HEIGHT)

    if section["image"]:
        height = 2.6 if section["table"] else BODY_HEIGHT
        _pptx_picture(slide, section["image"], right_left, BODY_TOP, 5.35, height)
    if section["table"]:
        top = BODY_TOP + 2.75 if section["image"] else BODY_TOP
        height = 1.7 if section["image"] else BODY_HEIGHT
        _pptx_table(slide, section["table"], right_left, top, 5.35, height)

    _pptx_footer(slide, deck_title, index, total)


def _pptx_body(slide, items: list, left: float, top: float, width: float, height: float) -> None:
    """要点区：## 小节标题、分级要点、普通段落分别排版。

    注意：这里的 left/top/width/height 都是**英寸浮点数**，必须经 ``Inches()``
    转成 EMU 再交给 python-pptx，否则会把小数写进 XML 造成文件损坏。
    """
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    first = True

    for kind, level, text in items:
        if not text:
            continue
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        size = PPTX_ITEM_SIZES.get(kind, 18)
        if kind == "sub":
            size = 21
        paragraph.space_after = Pt(9 if level == 0 else 5)
        paragraph.line_spacing = 1.2

        run = paragraph.add_run()
        run.text = text
        if kind == "sub":
            _pptx_run_font(
                run, PPTX_FONT, size, bold=True, color=PPTX_PALETTE["accent"]
            )
            _pptx_no_bullet(paragraph)
        elif kind == "bullet":
            _pptx_run_font(run, PPTX_FONT, size, color=PPTX_PALETTE["text"])
            _pptx_bullet(paragraph, PPTX_BULLETS[min(level, 2)], PPTX_PALETTE["accent"], level)
        elif kind == "quote":
            _pptx_run_font(run, PPTX_FONT, size, color=PPTX_PALETTE["muted"])
            _pptx_bullet(paragraph, "“", PPTX_PALETTE["muted"], level)
        elif kind == "code":
            _pptx_run_font(run, "Consolas", size, color=PPTX_PALETTE["text"])
            _pptx_no_bullet(paragraph)
        else:
            _pptx_run_font(run, PPTX_FONT, size, color=PPTX_PALETTE["text"])
            _pptx_no_bullet(paragraph)


def _pptx_picture(slide, source: str, left: float, top: float, width: float, height: float) -> None:
    """插图：等比缩放到给定框内，水平居中（参数均为英寸浮点数）。"""
    from pptx.util import Inches

    resolved = _resolve_image(source)
    if resolved is None:
        _pptx_text_box(
            slide, left, top, width, 0.4, f"［图片未找到：{source}］",
            size=12, color=PPTX_PALETTE["muted"],
        )
        return
    try:
        picture = slide.shapes.add_picture(
            str(resolved), Inches(left), Inches(top), width=Inches(width)
        )
    except Exception:      # noqa: BLE001 - 图片损坏时给文字占位
        _pptx_text_box(
            slide, left, top, width, 0.4, f"［图片无法插入：{resolved.name}］",
            size=12, color=PPTX_PALETTE["muted"],
        )
        return
    limit = int(Inches(height))
    if picture.height > limit:
        ratio = limit / picture.height
        picture.height = limit
        picture.width = int(picture.width * ratio)
    picture.left = int(Inches(left) + max(Inches(width) - picture.width, 0) / 2)


def _pptx_table(
    slide, rows: list[list[str]], left: float, top: float, width: float, height: float
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    total_rows = len(rows)
    total_cols = max((len(row) for row in rows), default=0)
    rows = [
        row[:PPTX_TABLE_MAX_COLS]
        for row in rows[:PPTX_TABLE_MAX_ROWS]
        if any(str(cell).strip() for cell in row)
    ]
    if not rows:
        return
    columns = max(len(row) for row in rows)
    shape = slide.shapes.add_table(
        len(rows), columns, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    table = shape.table
    for r, row in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                PPTX_PALETTE["accent"] if r == 0 else "FFFFFF"
            )
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    _pptx_run_font(
                        run,
                        PPTX_FONT,
                        13,
                        bold=(r == 0),
                        color=PPTX_PALETTE["cover_text"] if r == 0 else PPTX_PALETTE["text"],
                    )

    if total_rows > PPTX_TABLE_MAX_ROWS or total_cols > PPTX_TABLE_MAX_COLS:
        # 截断不能是「静默」的：否则模型与用户都以为整张表都上了页
        _pptx_text_box(
            slide,
            left,
            top + height + 0.05,
            width,
            0.3,
            f"（原表 {total_rows} 行 × {total_cols} 列，本页只显示前 "
            f"{PPTX_TABLE_MAX_ROWS} 行 / {PPTX_TABLE_MAX_COLS} 列）",
            size=11,
            color=PPTX_PALETTE["muted"],
        )


# ---------------------------------------------------------------- PDF
_CN_FONT = "MicrosoftYaHei"


def _register_cn_font() -> str:
    """注册中文字体，返回可用的字体名（失败时退回 reportlab 内置字体）。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if _CN_FONT in pdfmetrics.getRegisteredFontNames():
        return _CN_FONT
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont(_CN_FONT, str(candidate), subfontIndex=0))
                return _CN_FONT
            except Exception:        # pragma: no cover - 字体异常时退回
                continue
    return "Helvetica"


def cn_font_available() -> bool:
    """当前环境能否找到中文字体（找不到时 PDF 里的中文会变成方块）。"""
    return _register_cn_font() != "Helvetica"


def build_pdf(markdown: str, path: Path) -> None:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font = _register_cn_font()
    base = getSampleStyleSheet()
    styles = {
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName=font, fontSize=18, leading=24),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName=font, fontSize=15, leading=21),
        "h3": ParagraphStyle("h3", parent=base["Heading3"], fontName=font, fontSize=13, leading=19),
        "p": ParagraphStyle("p", parent=base["BodyText"], fontName=font, fontSize=11, leading=17,
                            alignment=TA_LEFT, spaceAfter=6),
        "li": ParagraphStyle("li", parent=base["BodyText"], fontName=font, fontSize=11, leading=17,
                             leftIndent=14, bulletIndent=4, spaceAfter=3),
        "quote": ParagraphStyle("quote", parent=base["BodyText"], fontName=font, fontSize=10,
                                leading=16, leftIndent=14, textColor="#5F5F58"),
        "code": ParagraphStyle("code", parent=base["Code"], fontName="Courier", fontSize=9,
                               leading=13, backColor="#F6F6F4", borderPadding=4, spaceAfter=8),
    }

    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title="Paper Agent 导出",
    )
    story: list = []
    for kind, text in parse_markdown_blocks(markdown):
        safe = _escape_xml(text)
        if kind == "code":
            safe = safe.replace("\n", "<br/>")
        if kind.startswith("li"):
            story.append(Paragraph(safe, styles["li"], bulletText="•"))
        elif kind in styles:
            story.append(Paragraph(safe, styles[kind]))
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(safe, styles["p"]))
            story.append(Spacer(1, 4))

    if not story:
        story.append(Paragraph("（空内容）", styles["p"]))
    doc.build(story)


def _escape_xml(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


# ---------------------------------------------------------------- 表格文件
def build_csv(content: str, path: Path) -> None:
    path.write_text(content if content.endswith("\n") else content + "\n", encoding="utf-8-sig")


def build_xlsx(content: str, path: Path, sheet_name: str = "Sheet1") -> None:
    """支持两种输入：CSV 文本，或 ``[[...], [...]]`` 形式的 JSON 二维数组。"""
    from openpyxl import Workbook

    rows: list[list] = []
    stripped = (content or "").strip()
    if stripped.startswith("["):
        try:
            data = json.loads(stripped)
            if isinstance(data, list):
                rows = [list(row) if isinstance(row, (list, tuple)) else [row] for row in data]
        except ValueError:
            rows = []
    if not rows:
        rows = [row for row in csv.reader(io.StringIO(content or ""))]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name or "Sheet1")[:31]
    for row in rows:
        sheet.append(row)
    workbook.save(str(path))


# ---------------------------------------------------------------- 读取
MAX_READ_BYTES = 25 * 1024 * 1024      # 单文件读取上限，避免把内存吃光


def read_document(path: str, max_chars: int = 8000) -> str:
    """按扩展名读取文档文本，供智能体分析已有资料。"""
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    if file.is_dir():
        raise IsADirectoryError(f"这是一个目录而不是文件：{path}")
    size = file.stat().st_size
    if size > MAX_READ_BYTES:
        # 各个分支都是「整份读进内存再截断」，大文件会把内存吃光
        raise ValueError(
            f"文件过大（{human_readable_size(size)}，上限 "
            f"{human_readable_size(MAX_READ_BYTES)}），"
            "请改用 execute_python 分块读取，或先压缩 / 拆分文件"
        )
    suffix = file.suffix.lower()

    if suffix in {".txt", ".md", ".csv", ".json", ".tex", ".py"}:
        return file.read_text(encoding="utf-8", errors="ignore")[:max_chars]
    if suffix == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(file))
        parts: list[str] = []
        for element in document.element.body.iterchildren():
            if element.tag == qn("w:p"):
                text = Paragraph(element, document).text.strip()
                if text:
                    parts.append(text)
                continue
            if element.tag == qn("w:tbl"):
                # 表格内容也要读出来，否则模型看不见自己写的实验数据
                for row in Table(element, document).rows:
                    parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return "\n".join(parts)[:max_chars]
    if suffix == ".pptx":
        from pptx import Presentation
        texts = []
        for slide in Presentation(str(file)).slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n".join(texts)[:max_chars]
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(file))
        return "\n".join((page.extract_text() or "") for page in reader.pages)[:max_chars]
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        sheet = load_workbook(str(file), data_only=True).active
        return "\n".join(
            ",".join("" if c is None else str(c) for c in row)
            for row in sheet.iter_rows(values_only=True)
        )[:max_chars]
    return file.read_text(encoding="utf-8", errors="ignore")[:max_chars]


# ---------------------------------------------------------------- Tools
class _ArtifactTool(Tool):
    """产物类工具基类：知道本会话已有文件，同名时原地更新而不是堆副本。"""

    def __init__(self, session_files: list[dict] | None = None) -> None:
        self._session_paths = {
            str(item.get("path", "")) for item in (session_files or []) if item.get("path")
        }

    def _target(
        self,
        filename: str,
        default_ext: str,
        content: str = "",
        append: bool = False,
    ) -> Path:
        name = _safe_name(filename, default_ext, _content_name(content))
        if append:
            # 续写必须打到「同名那份」上：走 _unique_path 去重的话会另起一个 -1 文件，
            # 模型以为在追加，实际是一堆散落的副本。同名那份优先在本会话产物目录里找，
            # 升级前生成的旧文件则在会话清单里（那批还平铺在产物根目录）。
            existing = session_artifacts.find_named(
                name, self._session_paths, base=ARTIFACTS_DIR
            )
            if existing is not None:
                return existing
        if Path(name).suffix.lower() != default_ext:
            # 工具的输出格式是固定的：模型写错扩展名（如 create_docx 传 .pdf）
            # 时强行纠正，否则会生成一个扩展名与内容不符、打不开的文件
            name = f"{Path(name).stem}{default_ext}"
        return _unique_path(name, self._session_paths)


class CreateDocxTool(_ArtifactTool):
    name = "create_docx"
    description = (
        "把 Markdown 内容生成 Word 文档（.docx）。"
        "写学位论文 / 开题报告 / 文献综述正文时用 doc_type=thesis，"
        "会按论文规范排版（正文宋体小四、黑体标题、1.5 倍行距、首行缩进 2 字符、"
        "三线表、图表题注、可更新的目录域、页脚页码），"
        "正文里用「![图 3-1 说明](配图文件名)」可以嵌入已生成的配图；"
        "**题注规范**：每个表格上方必须单独写一行「表 章-序 名称」（如表 3-3 数据库测试表），"
        "每张配图下方单独写一行「图 章-序 名称」——表题在表上方、图题在图下方，"
        "缺了题注的表格 / 图片不会被自动编号；"
        "普通通知、说明、记录等用 doc_type=plain。"
        "**长文务必分段写**：单次 markdown 控制在 6000 字以内，"
        "第一次调用不带 append，之后每写一部分就带 append=true 续写同一个文件名，"
        "否则单次输出会被截断。"
    )

    def __init__(self, session_files: list[dict] | None = None) -> None:
        super().__init__(session_files)
        self.parameters = [
            ToolParam("filename", "string", "文件名，如 第一章绪论.docx"),
            ToolParam("markdown", "string", "Markdown 正文（支持标题/列表/表格/代码块）"),
            ToolParam(
                "doc_type",
                "string",
                "thesis=按论文排版；plain=普通文档；留空按内容自动判断",
                required=False,
                enum=["thesis", "plain", "auto"],
            ),
            ToolParam(
                "append",
                "string",
                "true=往同名文档尾部续写（长文分段写入时用）；false/留空=新建或覆盖",
                required=False,
                enum=["true", "false"],
            ),
        ]

    def run(
        self,
        filename: str = "",
        markdown: str = "",
        doc_type: str = "auto",
        append: str = "false",
        **kwargs,
    ) -> ToolResult:
        text = (markdown or "").strip()
        if not text:
            # 没有正文就别落盘：否则工作区里会多出一份打不开的空文档
            return ToolResult(
                success=False,
                error="markdown 正文为空，没有生成文件。请把完整正文写进 markdown 参数后重新调用。",
            )
        continuing = str(append).strip().lower() == "true"
        path = self._target(filename, ".docx", text, append=continuing)
        if continuing and not path.exists():
            return ToolResult(
                success=False,
                error=(
                    f"找不到要续写的文档：{path.name}。"
                    "请先不带 append 调用一次把文件建出来，再续写。"
                ),
            )
        if continuing and (doc_type or "auto") == "auto":
            # 续写时不能只看这一段的字数：比如「参考文献」那段很短且没标题，
            # 单看会被判成普通文档。把已有正文一起拿来判，排版才前后一致。
            style = resolve_doc_style("auto", f"{_doc_head_text(path)}\n{text}", path.name)
        else:
            style = resolve_doc_style(doc_type, text, path.name)
        build_docx(text, path, style=style, append=continuing)
        label = "论文排版（宋体正文 / 黑体标题）" if style == "thesis" else "普通文档排版"
        action = "已续写到" if continuing else "已生成"
        return ToolResult(
            content=f"{action} Word 文档 {path.name}（本次 {len(text)} 字 · {label}）",
            artifact_paths=[str(path)],
        )


class CreatePptxTool(_ArtifactTool):
    name = "create_pptx"
    description = (
        "把 Markdown 大纲生成 16:9 演示文稿（.pptx），自带封面页、目录页、主题配色与装饰、"
        "分级要点、图表排版和页脚页码。"
        "写法：`#` 为一页标题（首个 `#` 与讲题相近、或只有短句时作为封面）；"
        "`##` 是**页内小节标题，不新开一页**；`- ` 为要点，缩进两格降一级；"
        "Markdown 表格会排成配色表格；`![图 x-y 说明](配图文件名)` 独占一行会插图到页面右侧。"
        f"一页要点超过 {MAX_ITEMS_PER_SLIDE} 条会自动拆成「（续）」页，"
        "**续页不会重复计入目录**；"
        "**不要自己写「目录」章节**：主标题达到 5 个时系统自动生成目录页（只列一级标题），"
        "不足 5 个则没有目录页。需要把核心算法分多页展开时，把该页内容写足即可（自动分页），"
        "或拆成多个一级标题（会各自成为一条目录）。"
        "**背景可以多样**：backgrounds 传多张（本地路径或画面描述，按章节顺序轮换）；"
        "某节要单独换背景，就在该节标题下写一行 `[背景图: 文件名或描述]`；"
        "系统还会逐页更换背景处理方式（整页蒙版 / 左侧渐变 / 右侧竖条）与裁切焦点，"
        "所以即使只有一张图也不会每页都一样。"
        "缺素材时先用 web_search(kind=image) 检索、download_file 下载，再把路径传给 backgrounds；"
        "不传则由系统按讲题联网生成（横版、无水印，同描述命中缓存复用）。"
    )

    def __init__(
        self,
        session_files: list[dict] | None = None,
        image_generator: Any = None,
    ) -> None:
        super().__init__(session_files)
        self._image_generator = image_generator
        self._cancel: Any = None
        self.parameters = [
            ToolParam("filename", "string", "文件名，如 开题答辩.pptx"),
            ToolParam(
                "markdown",
                "string",
                "Markdown 大纲：# 页标题、## 页内小节、- 要点（缩进降级）、| 表格 |、"
                "![图](路径)；某节标题下可写 [背景图: 文件名或描述] 单独换背景",
            ),
            ToolParam(
                "background",
                "string",
                "单张背景：留空或 auto=按讲题自动生成；也可给本地路径或画面描述",
                required=False,
            ),
            ToolParam(
                "backgrounds",
                "array",
                f"多张背景（本地路径或画面描述），按章节轮流使用，最多 {MAX_BACKDROPS} 张；"
                "想让每章背景不同就传多个",
                required=False,
            ),
        ]

    def set_cancel(self, cancel: Any) -> None:
        """文生图较慢，支持 Esc 中止。"""
        self._cancel = cancel

    def run(
        self,
        filename: str = "",
        markdown: str = "",
        background: str = "auto",
        backgrounds: list[str] | None = None,
        **kwargs,
    ) -> ToolResult:
        text = (markdown or "").strip()
        if not text:
            return ToolResult(
                success=False,
                error="markdown 大纲为空，没有生成文件。请把大纲写进 markdown 参数后重新调用。",
            )
        path = self._target(filename, ".pptx", text)
        stem = path.stem

        resolved, notes = self._resolve_backdrops(
            self._background_specs(background, backgrounds), text, stem
        )
        text, marker_note = self._resolve_markers(text, stem)
        if marker_note:
            notes.append(marker_note)

        build_pptx(text, path, backgrounds=resolved)
        detail = "、".join(item for item in notes if item) or "使用主题色底"
        return ToolResult(
            content=(
                f"已生成 PPT 演示文稿（含封面、目录与内容页排版；"
                f"{len(resolved)} 张背景按章节轮换 + 逐页变换处理方式）· {detail}"
            ),
            artifact_paths=[str(path)],
        )

    # ------------------------------------------------------------------ 背景
    @staticmethod
    def _background_specs(background: str, backgrounds: list[str] | None) -> list[str]:
        """把 background / backgrounds 归一成一组「路径或描述」。"""
        specs: list[str] = []
        if isinstance(backgrounds, str):        # 模型偶尔会把数组写成换行字符串
            backgrounds = re.split(r"[\n;；]+", backgrounds)
        for item in [background, *(backgrounds or [])]:
            value = (item or "").strip()
            if value and value not in specs:
                specs.append(value)
        return specs

    def _generate(self, description: str, markdown: str, stem: str) -> tuple[str, str]:
        """按描述生成/复用一张背景图，返回 ``(本地路径, 说明)``；失败时路径为空。"""
        if is_cancelled(self._cancel):
            return "", "已取消背景图生成，改用主题色底"
        try:
            local, reason = acquire_background(
                self._image_generator, _deck_topic(markdown, stem), description,
                cancel=self._cancel,
            )
        except Exception as exc:      # noqa: BLE001 - 出图失败不该让 PPT 生成失败
            return "", f"背景图生成失败（{exc}），已改用主题色底"
        return local or "", reason

    def _resolve_backdrops(
        self, specs: list[str], markdown: str, stem: str
    ) -> tuple[list[str], list[str]]:
        """把背景描述解析成本地图片路径：本地文件直接用，描述则联网生成。"""
        resolved: list[str] = []
        notes: list[str] = []
        local_files = 0
        generated = 0
        file_capped = False
        wants_auto = not specs
        for spec in specs:
            if is_auto(spec):
                wants_auto = True
                continue
            if _resolve_image(spec) is not None:
                if local_files >= MAX_LOCAL_BACKDROPS:
                    file_capped = True
                    continue
                resolved.append(spec)
                local_files += 1
                continue
            if Path(spec).suffix.lower() in IMAGE_FILE_EXTENSIONS:
                notes.append(f"背景图不存在（{spec}），已改用主题色底")
                continue
            if generated >= MAX_BACKDROPS:
                notes.append(f"背景最多 {MAX_BACKDROPS} 张，多余的描述已忽略")
                break
            local, reason = self._generate(spec, markdown, stem)
            generated += 1
            if local:
                resolved.append(local)
            else:
                notes.append(reason or f"背景「{spec[:14]}」生成失败")

        if file_capped:
            notes.append(f"本地背景最多 {MAX_LOCAL_BACKDROPS} 张，多余的已忽略")
        if local_files:
            notes.append(
                "已套用指定的背景图"
                if local_files == 1 and not resolved[1:]
                else f"已套用指定的背景图（{local_files} 张轮换）"
            )

        if not resolved and wants_auto:
            # 兼容原行为：没给背景就按讲题自动生成一张
            local, reason = self._generate("", markdown, stem)
            if local:
                resolved.append(local)
                notes.append("已按讲题自动生成背景图")
            else:
                notes.append(reason or "未准备背景图，改用主题色底")
        return resolved, notes

    def _resolve_markers(self, markdown: str, stem: str) -> tuple[str, str]:
        """把 `[背景图: 描述]` 里的描述落成真实图片，标记替换为文件名。"""
        if "[背景图" not in markdown:
            return markdown, ""
        lines: list[str] = []
        generated = 0
        replaced = 0
        for line in markdown.split("\n"):
            marker = BACKGROUND_MARKER.match(line.strip())
            if marker is None:
                lines.append(line)
                continue
            value = marker.group(1).strip()
            if _resolve_image(value) is not None:
                lines.append(line)
                continue
            local = ""
            if generated < MAX_BACKDROPS:
                local = self._generate(value, markdown, stem)
                generated += 1
            if local:
                lines.append(f"[背景图: {Path(local).name}]")
                replaced += 1
            else:
                lines.append("")            # 生成失败：丢掉这条指令，别当正文渲染
        note = f"{replaced} 处章节背景已单独生成" if replaced else ""
        return "\n".join(lines), note


class CreatePdfTool(_ArtifactTool):
    name = "create_pdf"
    description = "把 Markdown 内容生成 PDF 文档，适合定稿导出。"

    def __init__(self, session_files: list[dict] | None = None) -> None:
        super().__init__(session_files)
        self.parameters = [
            ToolParam("filename", "string", "文件名，如 论文初稿.pdf"),
            ToolParam("markdown", "string", "Markdown 正文"),
        ]

    def run(self, filename: str = "", markdown: str = "", **kwargs) -> ToolResult:
        text = (markdown or "").strip()
        if not text:
            return ToolResult(
                success=False,
                error="markdown 正文为空，没有生成文件。请把完整正文写进 markdown 参数后重新调用。",
            )
        path = self._target(filename, ".pdf", text)
        build_pdf(text, path)
        note = ""
        if not cn_font_available():
            # 静默回退会让用户拿到一份中文全是方块的 PDF，必须说明
            note = (
                "；注意：本机未找到中文字体，PDF 中的中文可能显示为方块，"
                "建议改用 create_docx（Word 版中文排版更稳）"
            )
        return ToolResult(content=f"已生成 PDF 文档{note}", artifact_paths=[str(path)])


class CreateCsvTool(_ArtifactTool):
    name = "create_csv"
    description = "生成 CSV 表格文件（.csv），适合实验数据、对照表。"

    def __init__(self, session_files: list[dict] | None = None) -> None:
        super().__init__(session_files)
        self.parameters = [
            ToolParam("filename", "string", "文件名，如 实验数据.csv"),
            ToolParam("content", "string", "CSV 文本内容，第一行为表头"),
        ]

    def run(self, filename: str = "", content: str = "", **kwargs) -> ToolResult:
        text = (content or "").strip()
        if not text:
            return ToolResult(
                success=False,
                error="content 内容为空，没有生成文件。请把表格内容（首行为表头）写进 content 参数后重新调用。",
            )
        path = self._target(filename, ".csv", text)
        build_csv(text, path)
        return ToolResult(content="已生成 CSV 文件", artifact_paths=[str(path)])


class CreateXlsxTool(_ArtifactTool):
    name = "create_xlsx"
    description = "生成 Excel 工作簿（.xlsx），内容可为 CSV 文本或 JSON 二维数组。"

    def __init__(self, session_files: list[dict] | None = None) -> None:
        super().__init__(session_files)
        self.parameters = [
            ToolParam("filename", "string", "文件名，如 数据统计.xlsx"),
            ToolParam("content", "string", "CSV 文本或 JSON 二维数组"),
            ToolParam("sheet_name", "string", "工作表名", required=False),
        ]

    def run(self, filename: str = "", content: str = "", sheet_name: str = "Sheet1", **kwargs) -> ToolResult:
        text = (content or "").strip()
        if not text:
            return ToolResult(
                success=False,
                error="content 内容为空，没有生成文件。请把表格内容（CSV 文本或 JSON 数组）写进 content 参数后重新调用。",
            )
        path = self._target(filename, ".xlsx", text)
        build_xlsx(text, path, sheet_name or "Sheet1")
        return ToolResult(content="已生成 Excel 文件", artifact_paths=[str(path)])


class ReadDocumentTool(Tool):
    name = "read_document"
    description = "读取本地文档内容用于分析，支持 docx/pptx/pdf/xlsx/csv/txt/md 等。"

    def __init__(self) -> None:
        self.parameters = [
            ToolParam("path", "string", "文件绝对路径"),
        ]

    def run(self, path: str = "", **kwargs) -> ToolResult:
        text = read_document(path)
        return ToolResult(content=text or "（文档无文本内容）")


class ListArtifactsTool(Tool):
    name = "list_artifacts"
    description = (
        "列出本会话工作区已有的文件（含上传与各轮生成的产物）及其路径。"
        "一般不需要调用：这些文件已在系统提示里给出，直接用 read_document 读即可。"
    )

    def __init__(self, session_files: list[dict] | None = None) -> None:
        self.parameters = []
        self._session_files = session_files or []

    def run(self, **kwargs) -> ToolResult:
        if self._session_files:
            listing = "\n".join(format_workspace_file(item) for item in self._session_files)
            return ToolResult(
                content=f"本会话工作区文件（共 {len(self._session_files)} 个）：\n{listing}"
            )
        # 兜底：拿不到会话清单时才扫（只扫**本会话**目录 + 根目录里的旧文件，
        # 别的会话目录不进 —— 一起扫就会把别人的产物列进来）
        files = sorted(
            (Path(item) for item in session_artifacts.candidates(base=ARTIFACTS_DIR)),
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        listing = "\n".join(
            f"- {f.name}（{f.stat().st_size} 字节）路径：{f}" for f in files[:30]
        )
        return ToolResult(content=listing or "（暂无产物）")


def register_document_skills(
    registry: ToolRegistry,
    session_files: list[dict] | None = None,
    image_generator: Any = None,
    background_generator: Any = None,
) -> ToolRegistry:
    """把全部文档 Skills 注册进工具表。

    Args:
        session_files: 本会话工作区文件清单，用于让产物列举限定在本会话范围内
        image_generator: 文生图生成器
        background_generator: PPT 背景图专用生成器（横版、无水印）；
            不传则复用 ``image_generator``
    """
    for tool in (
        CreateDocxTool(session_files),
        CreatePptxTool(session_files, background_generator or image_generator),
        CreatePdfTool(session_files),
        CreateCsvTool(session_files),
        CreateXlsxTool(session_files),
        ReadDocumentTool(),
        ListArtifactsTool(session_files),
        ExecutePythonTool(),
    ):
        registry.register(tool)
    return registry


def register_code_mode_skills(
    registry: ToolRegistry,
    workspace: Path | None = None,
    image_generator=None,
    video_generator=None,
    video_images_provider=None,
    video_progress=None,
    image_images_provider=None,
    session_files: list[dict] | None = None,
) -> ToolRegistry:
    """编程模式的工具集：读写代码文件 / 列目录 / 跑命令 / 执行 Python（+ 文生图）。

    刻意**不注册** create_docx / create_pptx 等文档工具：编程模式下要的是
    工程文件与命令，混进文档工具只会让模型走错路。

    但**文生图必须注册**：写工程时同样会要图标、示意图、占位图，用户也可能直接说
    「生成一张图」。早先漏了这一个，模型翻遍工具列表都找不到生图能力，只能回答
    「我没有图像生成能力」——明明文生图是配好的。

    Args:
        workspace: 该会话专属的工程目录（不传则用共享默认目录）
        image_generator: 文生图生成器；不是 None 时注册 ``generate_image`` 工具
        session_files: 本会话工作区文件清单（用户上传的图 + 之前生成的产物）。用来把
            模型给的**文件名**解析成本地路径；不传的话，编程模式下按名字取素材（
            ``anchor`` / ``reference_image`` / ``first_frame`` / ``merge_videos``）
            会一律「找不到」（真实会话 ``c754edb8c3d0`` 就是这么撞上的）。
    """
    from paper_agent.services.skills.code_workspace import code_root, register_code_skills

    root = workspace or code_root()
    # 按名字找文件时还要能命中**会话工程目录**：用户上传的图会被复制进去、模型自己也常
    # 把素材写在那里，而它们都不在 data/artifacts 里。只给 session_files 还不够。
    extra_dirs = [str(root)]
    registry.register(ReadDocumentTool())
    registry.register(ExecutePythonTool(workspace=root, relaxed=True))
    register_code_skills(registry, root)
    if image_generator is not None:
        from paper_agent.services.skills.image_skills import GenerateImageTool

        registry.register(
            GenerateImageTool(
                image_generator,
                session_files,
                images_provider=image_images_provider,
                extra_dirs=extra_dirs,
            )
        )
    if video_generator is not None:
        from paper_agent.services.skills.video_skills import (
            GenerateVideoTool,
            MergeVideosTool,
        )

        registry.register(
            GenerateVideoTool(
                video_generator,
                video_images_provider,
                session_files,
                progress=video_progress,
                extra_dirs=extra_dirs,
            )
        )
        # 已有几段想接起来时用：不重新生成、不花额度（也免得模型自己去调 ffmpeg）
        registry.register(MergeVideosTool(session_files, extra_dirs))
    from paper_agent.services.web_tools import register_web_skills

    return register_web_skills(registry)


def build_default_registry(
    image_generator=None,
    session_files: list[dict] | None = None,
    background_generator=None,
    mode: str = "doc",
    workspace: Path | None = None,
    video_generator=None,
    video_images_provider=None,
    video_progress=None,
    image_images_provider=None,
) -> ToolRegistry:
    """构造工具注册表（按工作模式装配）。

    Args:
        image_generator: 文生图生成器；传入时注册 ``generate_image`` 工具。
        session_files: 本会话工作区文件清单。
        background_generator: PPT 背景图专用生成器（横版、无水印）。
        mode: ``doc`` 文档模式（默认，写论文/做 PPT）；``code`` 编程模式
            （配环境 / 写代码 / 跑代码，工具与沙箱策略都不同）。
        workspace: 编程模式下该会话专属的工程目录（会话之间互不干扰）。
        video_generator: 文生视频生成器；传入时注册 ``generate_video`` 工具。
        video_images_provider: 本轮的图片素材（图生视频用，返回图片地址列表）。
        video_progress: 视频进度桥；长视频分段生成时用它写「第 i/N 段」。
        image_images_provider: 本轮的图片素材（图生图用）；有人在说人物形象而模型
            没填 ``reference_image`` 时，工具会拿它兜底（否则会按文字另画一张脸）。
    """
    if mode == "code":
        return register_code_mode_skills(
            ToolRegistry(),
            workspace,
            image_generator,
            video_generator,
            video_images_provider,
            video_progress,
            image_images_provider,
            session_files=session_files,
        )

    registry = register_document_skills(
        ToolRegistry(), session_files, image_generator, background_generator
    )
    if image_generator is not None:
        from paper_agent.services.skills.image_skills import GenerateImageTool

        # session_files：把 reference_image 里的文件名解析成本地图（图生图用）
        registry.register(
            GenerateImageTool(
                image_generator, session_files, images_provider=image_images_provider
            )
        )
    if video_generator is not None:
        from paper_agent.services.skills.video_skills import (
            GenerateVideoTool,
            MergeVideosTool,
        )

        # session_files：把 continue_from / 首尾帧里的文件名解析成本地文件
        registry.register(
            GenerateVideoTool(
                video_generator,
                video_images_provider,
                session_files,
                progress=video_progress,
            )
        )
        # 已有几段想接起来时用：不重新生成、不花额度（也免得模型自己去调 ffmpeg）
        registry.register(MergeVideosTool(session_files))
    # 联网检索与素材下载：找资料、找背景图用（免密，下载有安全边界）
    from paper_agent.services.web_tools import register_web_skills

    return register_web_skills(registry)
