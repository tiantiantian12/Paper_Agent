"""文档预览：把产物文件转成聊天内可直接渲染的内容。

纯逻辑模块（不依赖 Qt），便于单测；界面层见 ``ui/chat/document_preview.py``。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from paper_agent.core.constants import IMAGE_FILE_EXTENSIONS

PREVIEW_LIMIT = 4000        # 预览的最大字符数
MAX_TABLE_ROWS = 40         # 表格最多预览多少行
MAX_TABLE_COLUMNS = 12

MARKDOWN_SUFFIXES = {".md", ".markdown"}
TEXT_SUFFIXES = {
    ".txt", ".log", ".ini", ".cfg", ".conf", ".yaml", ".yml", ".toml",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs",
    ".json", ".xml", ".html", ".css", ".sql", ".sh", ".bat", ".ps1", ".tex",
}
_HEADING_STYLE = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)


@dataclass
class Preview:
    """预览结果。

    Attributes:
        kind: ``markdown`` / ``text`` / ``image`` / ``unavailable``
        content: Markdown、纯文本，或图片路径
        note: 顶部提示（不支持、失败原因等）；为空表示正常
        truncated: 内容是否被截断
    """

    kind: str = "unavailable"
    content: str = ""
    note: str = ""
    truncated: bool = False

    @property
    def available(self) -> bool:
        return self.kind != "unavailable"


def build_preview(path: str | Path, limit: int = PREVIEW_LIMIT) -> Preview:
    """按文件类型生成预览内容；任何异常都收敛成「不可预览 + 原因」。"""
    target = Path(path)
    if not target.exists():
        return Preview(note="文件不存在，可能已被清理")

    suffix = target.suffix.lower()
    if suffix in IMAGE_FILE_EXTENSIONS:
        return Preview(kind="image", content=str(target))

    try:
        if suffix == ".docx":
            text = _docx_markdown(target, limit)
        elif suffix == ".pdf":
            text = _pdf_markdown(target, limit)
        elif suffix == ".pptx":
            text = _pptx_markdown(target, limit)
        elif suffix in {".csv", ".tsv"}:
            text = _table_markdown(_csv_rows(target))
        elif suffix in {".xlsx", ".xlsm"}:
            text = _table_markdown(_xlsx_rows(target))
        elif suffix in MARKDOWN_SUFFIXES:
            return _clip_preview(_read_text(target), limit, "markdown")
        elif suffix in TEXT_SUFFIXES:
            return _clip_preview(_read_text(target), limit, "text")
        else:
            return Preview(
                note=f"暂不支持预览 {suffix or '该类型'} 文件，点卡片可用系统默认程序打开"
            )
    except Exception as exc:      # noqa: BLE001 - 预览失败不该影响聊天
        return Preview(note=f"预览失败：{exc}")

    return _clip_preview(text, limit, "markdown")


# ---------------------------------------------------------------- 截断
def _clip_preview(text: str, limit: int, kind: str) -> Preview:
    text = (text or "").strip()
    if not text:
        return Preview(kind=kind, note="文件内容为空")
    if len(text) <= limit:
        return Preview(kind=kind, content=text)
    return Preview(
        kind=kind,
        content=f"{text[:limit].rstrip()}\n\n……（全文约 {len(text)} 字，此处仅预览前 {limit} 字）",
        truncated=True,
    )


def _read_text(path: Path) -> str:
    # utf-8-sig 会顺带吃掉 BOM：生成的 CSV 带 BOM，否则第一个表头会多出不可见字符
    return path.read_text(encoding="utf-8-sig", errors="ignore")


# ---------------------------------------------------------------- Word
def _docx_markdown(path: Path, limit: int) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    parts: list[str] = []
    length = 0
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _paragraph_level(paragraph)
            parts.append(f"{'#' * level} {text}" if level else text)
        elif element.tag == qn("w:tbl"):
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in Table(element, document).rows
            ]
            table = _table_markdown(rows)
            if table:
                parts.append(table)
        length = sum(len(part) for part in parts)
        if length > limit * 1.5:
            break

    shapes = len(document.inline_shapes)
    if shapes:
        parts.append(f"（文档内嵌 {shapes} 张图片，预览不显示图片）")
    return "\n\n".join(parts)


def _paragraph_level(paragraph) -> int:
    name = (getattr(paragraph.style, "name", "") or "").strip()
    if name.lower() in {"title", "标题"}:
        return 1
    match = _HEADING_STYLE.search(name)
    return min(int(match.group(1)), 6) if match else 0


# ---------------------------------------------------------------- PDF / PPT
def _pdf_markdown(path: Path, limit: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        title = f"## 第 {index} 页" if len(reader.pages) > 1 else ""
        pages.append(f"{title}\n\n{text}" if title else text)
        if sum(len(item) for item in pages) > limit * 1.5:
            pages.append(f"……（共 {len(reader.pages)} 页，此处仅预览前 {index} 页）")
            break
    return "\n\n".join(pages)


def _pptx_markdown(path: Path, limit: int) -> str:
    from pptx import Presentation

    slides: list[str] = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if not texts:
            continue
        slides.append(f"## 第 {index} 页\n\n" + "\n\n".join(texts))
        if sum(len(item) for item in slides) > limit * 1.5:
            break
    return "\n\n".join(slides)


# ---------------------------------------------------------------- 表格类
def _csv_rows(path: Path) -> list[list[str]]:
    import csv
    import io

    content = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(content[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(content), dialect)]


def _xlsx_rows(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    sheet = load_workbook(str(path), data_only=True).active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell) for cell in row])
        if len(rows) > MAX_TABLE_ROWS:
            break
    return rows


def _table_markdown(rows: list[list[str]]) -> str:
    rows = [
        row[:MAX_TABLE_COLUMNS]
        for row in rows[:MAX_TABLE_ROWS]
        if any(str(cell).strip() for cell in row)
    ]
    if not rows:
        return ""
    columns = max(len(row) for row in rows)
    lines = [
        "| " + " | ".join(_escape_cell(row, columns)) + " |"
        for row in rows
    ]
    lines.insert(1, "| " + " | ".join(["---"] * columns) + " |")
    if len(rows) >= MAX_TABLE_ROWS:
        lines.append(f"\n（仅预览前 {MAX_TABLE_ROWS} 行）")
    return "\n".join(lines)


def _escape_cell(row: list[str], columns: int) -> list[str]:
    padded = [str(cell).replace("|", "\\|").replace("\n", " ") for cell in row]
    padded += [""] * (columns - len(padded))
    return padded
