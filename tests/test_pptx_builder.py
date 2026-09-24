"""PPT 生成：封面、目录、版式、分级要点、图表、分页与背景图。"""

from __future__ import annotations

import re

import pytest
from pptx import Presentation
from pptx.util import Inches

from paper_agent.services.skills.document_skills import build_pptx

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

DECK = """# 基于YOLOv8的车牌识别检测算法研究

答辩人：张三

# 研究背景

- 车牌识别是智能交通的核心环节
  - 传统方法依赖手工特征
- 复杂场景下小目标检测仍存在挑战

# 算法改进

| 模型 | mAP |
| --- | --- |
| YOLOv8n | 0.874 |
| AC-YOLOv8 | 0.912 |

# 结论与展望

- 完成了轻量化检测模型
"""


def _deck(tmp_path, markdown: str = DECK, name: str = "答辩.pptx", **kwargs):
    path = tmp_path / name
    build_pptx(markdown, path, **kwargs)
    return Presentation(str(path))


def _slide_texts(slide) -> list[str]:
    return [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


def _all_texts(deck) -> str:
    return "\n".join("\n".join(_slide_texts(slide)) for slide in deck.slides)


# ---------------------------------------------------------------- 基本结构
def test_slide_size_is_16_9(tmp_path):
    deck = _deck(tmp_path)
    assert deck.slide_width == Inches(13.333)
    assert deck.slide_height == Inches(7.5)


def test_all_shapes_have_valid_geometry(tmp_path, sample_png):
    """坐标必须是合法 EMU 整数：传浮点英寸会把手写坏，且只有读坐标时才暴露。"""
    markdown = (
        "# 讲题\n\n# 有图有表\n\n- 要点一\n\n"
        f"![图 1-1 架构]({sample_png})\n\n| 列 | 值 |\n| --- | --- |\n| a | 1 |\n"
    )
    deck = _deck(tmp_path, markdown)
    width, height = Inches(13.333), Inches(7.5)
    for index, slide in enumerate(deck.slides, start=1):
        assert slide.shapes, f"第 {index} 页没有内容"
        for shape in slide.shapes:
            assert isinstance(int(shape.left), int), f"第 {index} 页 left 非法"
            assert isinstance(int(shape.top), int), f"第 {index} 页 top 非法"
            assert shape.width > 0 and shape.height > 0
            assert 0 <= shape.left and shape.left + shape.width <= width, (
                f"第 {index} 页元素横向越界"
            )
            assert 0 <= shape.top and shape.top + shape.height <= height, (
                f"第 {index} 页元素纵向越界"
            )


def test_cover_and_sections(tmp_path):
    deck = _deck(tmp_path)
    # 封面 + 研究背景 + 算法改进 + 结论与展望
    assert len(deck.slides) == 4
    cover = _slide_texts(deck.slides[0])
    assert "基于YOLOv8的车牌识别检测算法研究" in cover
    assert "答辩人：张三" in cover, "封面段落的短句应作为副标题"
    assert "研究背景" in _slide_texts(deck.slides[1])


def test_cover_uses_gradient_background(tmp_path):
    deck = _deck(tmp_path)
    cover_xml = deck.slides[0]._element.xml
    assert "gradFill" in cover_xml, "封面应有主题色渐变底"
    assert "<a:alpha" in cover_xml, "封面装饰块需要半透明"


def test_content_page_has_title_bar_and_footer(tmp_path):
    deck = _deck(tmp_path)
    slide = deck.slides[1]
    assert slide.background.fill.type is not None
    texts = " ".join(_slide_texts(slide))
    assert "1 / 3" in texts, f"内容页要有页码：{texts}"
    assert "基于YOLOv8的车牌识别检测算法研究" in texts, "页脚应带讲题名"


def test_brand_stripe_on_content_pages(tmp_path):
    deck = _deck(tmp_path)
    # 内容页左侧主题色竖条 + 右下淡色装饰块
    shapes = deck.slides[1].shapes
    fills = [shape.fill.fore_color.rgb.__str__() for shape in shapes if shape.shape_type is not None and shape.fill.type is not None and shape.fill.type == 1]
    assert "1F4E79" in fills, f"应有主题色（accent）形状：{fills}"


# ---------------------------------------------------------------- 目录页
def test_toc_slide_for_many_sections(tmp_path):
    markdown = "# 答辩\n\n# 一\n\n- a\n\n# 二\n\n- b\n\n# 三\n\n- c\n\n# 四\n\n- d\n\n# 五\n\n- e\n"
    deck = _deck(tmp_path, markdown)
    assert "目录" in _slide_texts(deck.slides[1]), "章节 ≥5 应自动生成目录页"
    assert _toc_entries(deck.slides[1]) == ["一", "二", "三", "四", "五"]


def test_no_toc_for_short_deck(tmp_path):
    deck = _deck(tmp_path)
    assert "目录" not in _all_texts(deck)


def _toc_entries(slide) -> list[str]:
    """目录页的章节条目（排除标题、页脚与「还有 N 个章节」提示）。

    条目不显示页码前缀——用户明确反馈过「02 一、研究背景与意义」这种写法要去掉。
    """
    entries: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        if shape.top is None or shape.top > Inches(6.5):      # 页脚（讲题名 / 页码）
            continue
        stripped = shape.text_frame.text.strip()
        if stripped == "目录" or stripped.startswith("（还有"):
            continue
        entries.append(stripped)
    return entries


def test_toc_lists_sections_not_continuation_pages(tmp_path):
    """目录按一级标题算：自动分页出来的（续）页不能再各占一条。"""
    bullets = "\n".join(f"- 要点{i}" for i in range(1, 13))       # 12 条 → 拆 2 页
    markdown = (
        "# 讲题\n\n# 一、背景\n\n- a\n\n"
        f"# 二、核心算法\n\n{bullets}\n\n"
        "# 三、实验\n\n- c\n\n# 四、部署\n\n- d\n\n# 五、结论\n\n- e\n"
    )
    deck = _deck(tmp_path, markdown, name="答辩.pptx")

    toc = next(slide for slide in deck.slides if _slide_texts(slide)[0].strip() == "目录")
    entries = _toc_entries(toc)
    assert len(entries) == 5, f"5 个主标题应只列 5 条：{entries}"
    assert not any("（续）" in entry for entry in entries)

    titles = [_slide_texts(slide)[0].strip() for slide in deck.slides]
    assert "二、核心算法（续）" in titles, "内容仍要分页展开"
    assert len(deck.slides) == 1 + 1 + 6, "封面 + 目录 + 6 个内容页"


def test_handwritten_toc_section_is_ignored(tmp_path):
    """自己写「# 目录」既不该占一页，也不该出现在目录里。"""
    markdown = (
        "# 讲题\n\n"
        "# 目录\n\n- 一、背景\n- 二、算法\n\n"          # 手写目录章节
        "# 一、背景\n\n- a\n\n# 二、算法\n\n- b\n\n"
        "# 三、实验\n\n- c\n\n# 四、部署\n\n- d\n\n# 五、结论\n\n- e\n"
    )
    deck = _deck(tmp_path, markdown, name="答辩.pptx")

    titles = [_slide_texts(slide)[0].strip() for slide in deck.slides]
    assert titles.count("目录") == 1, f"只有系统生成的目录页：{titles}"
    toc = next(slide for slide in deck.slides if _slide_texts(slide)[0].strip() == "目录")
    entries = _toc_entries(toc)
    assert len(entries) == 5 and not any("目录" in entry for entry in entries), entries
    assert "手写的目录条目不该落成正文页" not in titles


def test_toc_threshold_counts_sections_not_pages(tmp_path):
    """章节不足 5 个时不出目录页，即使内容被拆成很多页。"""
    bullets = "\n".join(f"- 要点{i}" for i in range(1, 20))       # 19 条 → 拆 3 页
    markdown = f"# 讲题\n\n# 一、算法\n\n{bullets}\n\n# 二、实验\n\n- b\n\n# 三、结论\n\n- c\n"
    deck = _deck(tmp_path, markdown, name="答辩.pptx")

    assert len(deck.slides) == 6, "封面 + 算法 3 页 + 实验 + 结论"
    assert "目录" not in _all_texts(deck)


def test_deck_title_becomes_cover_not_toc_entry(tmp_path):
    """首个标题与讲题相近（其实是同一篇）时按封面处理，不重复出页、不进目录。"""
    markdown = (
        "# 基于YOLOv8的车牌识别检测算法研究\n\n答辩人：张三\n\n"   # 与文件名相近
        "# 一、背景\n\n- a\n\n# 二、算法\n\n- b\n\n"
        "# 三、实验\n\n- c\n\n# 四、部署\n\n- d\n\n# 五、结论\n\n- e\n"
    )
    deck = _deck(tmp_path, markdown, name="基于YOLOv8的车牌识别检测算法答辩PPT_v3.pptx")

    assert _slide_texts(deck.slides[0])[0].strip() == "基于YOLOv8的车牌识别检测算法研究"
    toc = next(slide for slide in deck.slides if _slide_texts(slide)[0].strip() == "目录")
    entries = _toc_entries(toc)
    assert len(entries) == 5, entries
    assert not any("车牌识别检测算法研究" in entry for entry in entries)


def test_version_suffix_stripped_from_filename(tmp_path):
    """文件名带 _v3 这类版本后缀时，封面标题不该带进去。"""
    deck = _deck(tmp_path, "# 一、背景\n\n- a\n", name="开题答辩报告_v3.pptx")
    assert _slide_texts(deck.slides[0])[0].strip() == "开题答辩报告"


# ---------------------------------------------------------------- 要点排版
def test_nested_bullets_have_levels_and_chars(tmp_path):
    deck = _deck(tmp_path)
    chars = [element.get("char") for element in deck.slides[1]._element.iter(A + "buChar")]
    assert "●" in chars and "○" in chars, f"一级/二级要点要用不同符号：{chars}"

    margins = {
        int(element.get("marL"))
        for element in deck.slides[1]._element.iter(A + "pPr")
        if element.get("marL")
    }
    assert len(margins) >= 2, f"二级要点缩进应更大：{margins}"


def test_sub_heading_and_body_have_no_bullet(tmp_path):
    deck = _deck(tmp_path, "# 章节\n\n## 页内小节\n\n普通说明文字。\n\n- 要点一条\n")
    xml = deck.slides[1]._element.xml
    assert "buNone" in xml, "页内小节标题与普通段落不该带项目符号"
    assert "▪" not in xml


def test_long_section_is_split(tmp_path):
    bullets = "\n".join(f"- 要点{i}" for i in range(1, 13))
    markdown = f"# 长章节\n\n{bullets}\n"
    deck = _deck(tmp_path, markdown)
    texts = [_slide_texts(slide) for slide in deck.slides]
    titles = [next((t for t in item if t.startswith("长章节")), "") for item in texts]
    assert any("（续）" in title for title in titles), f"超出应自动分页：{titles}"
    assert len(deck.slides) == 3, "封面 + 两页要点"


def test_east_asian_font_is_set(tmp_path):
    deck = _deck(tmp_path)
    typefaces = [
        element.get("typeface")
        for slide in deck.slides
        for element in slide._element.iter(A + "ea")
    ]
    assert typefaces and set(typefaces) == {"微软雅黑"}, f"中文必须写 a:ea：{set(typefaces)}"


# ---------------------------------------------------------------- 图表
def test_table_is_rendered(tmp_path):
    deck = _deck(tmp_path)
    slide = deck.slides[2]
    tables = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
    assert tables, "Markdown 表格应排成 PPT 表格"
    table = tables[0].table
    assert table.cell(0, 0).text == "模型"
    assert table.cell(1, 1).text == "0.874"
    assert table.cell(0, 0).fill.fore_color.rgb.__str__() == "1F4E79", "表头要用主题色"


def test_picture_is_inserted(tmp_path, sample_png):
    markdown = f"# 系统架构\n\n![图 1-1 系统架构]({sample_png})\n\n- 说明文字\n"
    deck = _deck(tmp_path, markdown)
    pictures = [
        shape for shape in deck.slides[1].shapes if shape.shape_type == 13
    ]  # 13 = PICTURE
    assert pictures, "图片块应插到页面上"
    assert pictures[0].width <= Inches(5.35) + 1


def test_missing_picture_shows_placeholder(tmp_path):
    deck = _deck(tmp_path, "# 图\n\n![图 1-1 缺失](不存在的图.png)\n")
    assert "图片未找到" in _all_texts(deck)


# ---------------------------------------------------------------- 背景图
def test_background_image_is_used_on_every_slide(tmp_path, sample_png):
    """每页都要有背景图（版式可不同：整页 / 左侧渐变 / 右侧竖条）。"""
    deck = _deck(tmp_path, background=str(sample_png))
    for slide in deck.slides:
        first = slide.shapes[0]
        assert first.shape_type == 13, "背景图应是每页第一个（最底层）形状"
        if first.width == Inches(13.333):
            # 整页铺图时必须压一层半透明蒙版，否则文字压在图上看不清
            assert "<a:alpha" in slide._element.xml, "整页铺图要压蒙版保证可读"

    # 封面与目录保持整页铺满；正文页允许换成别的版式
    assert deck.slides[0].shapes[0].width == Inches(13.333)
    assert deck.slides[1].shapes[0].width == Inches(13.333)
    configs = {
        (round(slide.shapes[0].left / 914400, 2), round(slide.shapes[0].width / 914400, 2))
        for slide in list(deck.slides)[2:]
    }
    assert len(configs) >= 2, f"正文页不该每页都用同一种背景版式：{configs}"


def test_background_missing_falls_back_to_theme(tmp_path):
    deck = _deck(tmp_path, background="不存在的背景.png")
    assert "gradFill" in deck.slides[0]._element.xml
    assert deck.slides[0].shapes[0].shape_type != 13


# ---------------------------------------------------------------- 兜底
def test_empty_markdown_still_builds(tmp_path):
    deck = _deck(tmp_path, "")
    assert len(deck.slides) >= 1


def test_single_section_becomes_cover_and_content(tmp_path):
    """只有一个 # 时：没要点就只做封面，有要点则合成封面 + 内容页。"""
    cover_only = _deck(tmp_path, "# 开题答辩\n\n汇报人：张三\n", name="a.pptx")
    assert len(cover_only.slides) == 1
    assert "开题答辩" in _slide_texts(cover_only.slides[0])

    with_points = _deck(tmp_path, "# 研究背景\n\n- 现状\n- 问题\n", name="b.pptx")
    assert len(with_points.slides) == 2
    assert any("研究背景" in text for text in _slide_texts(with_points.slides[1]))


# ---------------------------------------------------------------- 工具
def test_tool_writes_file(tmp_path, artifacts_dir):
    from paper_agent.services.skills.document_skills import CreatePptxTool

    tool = CreatePptxTool()
    result = tool.run(filename="开题答辩.pptx", markdown=DECK)
    assert result.artifact_paths[0].endswith("开题答辩.pptx")
    assert "封面" in result.content
    assert Presentation(result.artifact_paths[0]).slides

    with_background = tool.run(
        filename="带背景.pptx", markdown=DECK, background="不存在的图.png"
    )
    assert "背景图" in with_background.content


@pytest.mark.parametrize("name", ["答辩", "答辩.pptx", "答辩.pdf"])
def test_tool_forces_pptx_extension(tmp_path, artifacts_dir, name):
    from paper_agent.services.skills.document_skills import CreatePptxTool

    result = CreatePptxTool().run(filename=name, markdown="# 标题\n\n- 要点")
    assert result.artifact_paths[0].endswith(".pptx")
