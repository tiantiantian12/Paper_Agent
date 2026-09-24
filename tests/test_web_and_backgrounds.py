"""联网检索 / 下载，以及 PPT 多背景与逐页变体的回归用例（不依赖网络）。"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation

from paper_agent.services import web_tools
from paper_agent.services.skills.document_skills import (
    BACKGROUND_MARKER,
    MAX_BACKDROPS,
    build_pptx,
)
from paper_agent.services.web_tools import download_file, is_public_url, search_images


# ---------------------------------------------------------------- 检索解析（离线样本）
BING_HTML = (
    '<ol id="b_results">'
    '<li class="b_algo" data-id iid=SERP.1><link rel="stylesheet" href="/x.css"/>'
    '<h2><a href="https://example.com/a" h="ID=SERP,1">第一篇 结果</a></h2>'
    '<p class="b_lineclamp2">这是第一篇的摘要内容</p></li>'
    '<li class="b_algo" data-id iid=SERP.2><h2><a href="https://example.com/b">第二篇</a></h2>'
    "<p>第二篇摘要</p></li>"
    "</ol>"
)
SO_HTML = (
    '<div class="res-list"><h3><a href="https://so.example.com/1">360 结果一</a></h3>'
    '<p class="res-desc">360 的摘要</p></div>'
    '<div class="res-list"><h3><a href="https://so.example.com/2">360 结果二</a></h3></div>'
)
IMG_HTML = (
    '{"murl&quot;:&quot;https://img.example.com/a.jpg&quot;,'
    '&quot;t&quot;:&quot;蓝色科技背景&quot;},'
    '{"murl&quot;:&quot;https://img.example.com/b.png&quot;}'
)
BING_IMG_HTML = '&quot;murl&quot;:&quot;https://img.example.com/a.jpg&quot;,&quot;t&quot;:&quot;蓝图&quot;'


def test_bing_result_parsing_without_network(monkeypatch):
    monkeypatch.setattr(web_tools, "_fetch", lambda url, timeout=0, referer="": BING_HTML.encode())
    results = web_tools._search_bing("任意关键词", 5)
    assert [item["url"] for item in results] == [
        "https://example.com/a",
        "https://example.com/b",
    ]
    assert results[0]["title"] == "第一篇 结果"
    assert "摘要" in results[0]["snippet"]


def test_360_result_parsing_without_network(monkeypatch):
    monkeypatch.setattr(web_tools, "_fetch", lambda url, timeout=0, referer="": SO_HTML.encode())
    results = web_tools._search_360("任意关键词", 5)
    assert [item["url"] for item in results] == [
        "https://so.example.com/1",
        "https://so.example.com/2",
    ]
    assert results[0]["snippet"] == "360 的摘要"


def test_bing_image_parsing_without_network(monkeypatch):
    monkeypatch.setattr(
        web_tools, "_fetch", lambda url, timeout=0, referer="": BING_IMG_HTML.encode()
    )
    results = web_tools._search_bing_images("科技感背景", 4)
    assert results[0]["url"] == "https://img.example.com/a.jpg"
    assert results[0]["title"] == "蓝图"


def test_search_falls_back_to_next_provider(monkeypatch):
    """必应挂了要自动降级到 360，而不是直接失败。"""
    calls = {"n": 0}

    def boom(query, count):
        calls["n"] += 1
        raise RuntimeError("必应超时")

    monkeypatch.setattr(web_tools, "_search_bing", boom)
    monkeypatch.setattr(
        web_tools, "_search_360", lambda query, count: [{"title": "t", "url": "u", "snippet": ""}]
    )
    results = web_tools.search_web("关键词", 3)
    assert calls["n"] == 1 and results[0]["url"] == "u"


def test_all_providers_failing_gives_readable_error(monkeypatch):
    for name in ("_search_bing", "_search_360", "_search_duckduckgo"):
        monkeypatch.setattr(
            web_tools, name, lambda query, count: (_ for _ in ()).throw(RuntimeError("超时"))
        )
    with pytest.raises(RuntimeError) as info:
        web_tools.search_web("关键词", 3)
    assert "所有检索源都失败" in str(info.value)


# ---------------------------------------------------------------- 下载安全
@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1:8080/x.png",
        "http://localhost/x.png",
        "http://192.168.1.10/admin.png",
        "http://10.0.0.5/a.png",
        "http://169.254.169.254/latest/meta-data",
        "file:///C:/Windows/win.ini",
        "ftp://example.com/a.png",
        "",
    ],
)
def test_private_or_unsupported_urls_are_rejected(url):
    assert is_public_url(url) is False


def test_download_rejects_private_url(tmp_path):
    with pytest.raises(ValueError) as info:
        download_file("http://127.0.0.1/x.png", "x.png")
    assert "只允许公网" in str(info.value)


def test_safe_filename_uses_real_suffix():
    """用户给了 .png 但内容其实是 jpeg 时，不能生成 x.png.jpg。"""
    assert web_tools._safe_filename("背景.png", ".jpg") == "背景.jpg"
    assert web_tools._safe_filename("背景", ".jpg") == "背景.jpg"
    assert web_tools._safe_filename("a/b:c*?.png", ".png") == "b_c__.png"
    assert web_tools._safe_filename("", ".png").endswith(".png")


class _Headers:
    def __init__(self, content_type: str) -> None:
        self._content_type = content_type

    def get_content_type(self) -> str:
        return self._content_type


class _FakeResponse(io.BytesIO):
    def __init__(self, payload: bytes, content_type: str) -> None:
        super().__init__(payload)
        self.headers = _Headers(content_type)

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


def test_download_writes_real_image(tmp_path, monkeypatch):
    """用假响应验证落盘流程（不联网）。"""
    payload = b"\x89PNG\r\n\x1a\n" + b"0" * 200
    monkeypatch.setattr(web_tools, "is_public_url", lambda url: True)
    monkeypatch.setattr(
        web_tools.urllib.request,
        "urlopen",
        lambda request, timeout=None: _FakeResponse(payload, "image/png"),
    )
    path = download_file("https://img.example.com/a.png", "素材.png", directory=tmp_path)
    assert path.name == "素材.png"
    assert path.read_bytes() == payload


def test_download_rejects_non_image(tmp_path, monkeypatch):
    monkeypatch.setattr(web_tools, "is_public_url", lambda url: True)
    monkeypatch.setattr(
        web_tools.urllib.request,
        "urlopen",
        lambda request, timeout=None: _FakeResponse(b"<html></html>", "text/html"),
    )
    with pytest.raises(ValueError) as info:
        download_file("https://example.com/page", "x.html", directory=tmp_path)
    assert "不支持的类型" in str(info.value)


def test_tools_are_registered():
    from paper_agent.services.skills import build_default_registry

    registry = build_default_registry()
    names = {spec["function"]["name"] for spec in registry.specs()}
    assert {"web_search", "download_file"} <= names


# ---------------------------------------------------------------- PPT 背景
def _picture_ids(slide) -> list[str]:
    return [
        shape.image.sha1
        for shape in slide.shapes
        if shape.shape_type == 13 and hasattr(shape, "image")
    ]


def _fake_images(tmp_path) -> list[str]:
    from PIL import Image

    paths = []
    for index, color in enumerate(((20, 40, 90), (120, 30, 60)), start=1):
        target = tmp_path / f"bg{index}.png"
        Image.new("RGB", (1600, 900), color).save(target)
        paths.append(str(target))
    return paths


SECTIONS = (
    "# 讲题\n\n# 一、背景\n\n- a\n\n# 二、现状\n\n- b\n\n"
    "# 三、算法\n\n- c\n\n# 四、实验\n\n- d\n\n# 五、结论\n\n- e\n"
)


def test_marker_is_parsed_as_background_not_text(tmp_path):
    markdown = SECTIONS.replace("# 三、算法\n", "# 三、算法\n[背景图: bg1.png]\n")
    from paper_agent.services.skills.document_skills import _pptx_sections, parse_markdown_blocks

    sections, _ = _pptx_sections(parse_markdown_blocks(markdown))
    target = next(item for item in sections if item["title"] == "三、算法")
    assert target["background"] == "bg1.png"
    assert all("背景图" not in text for _kind, _level, text in target["items"])
    assert BACKGROUND_MARKER.match("[背景图: 深蓝科技感]").group(1) == "深蓝科技感"


def test_multiple_backdrops_rotate_by_chapter(tmp_path):
    images = _fake_images(tmp_path)
    path = tmp_path / "轮换.pptx"
    build_pptx(SECTIONS, path, backgrounds=images)
    deck = Presentation(str(path))

    content = []
    for slide in list(deck.slides)[2:]:        # 跳过封面与目录
        ids = _picture_ids(slide)
        if ids:
            content.append(ids[0])
    assert len(set(content)) == len(images), f"两张背景应该在章节间交替：{content}"


def test_pages_are_not_all_identical(tmp_path):
    """只有一张背景图时，逐页处理方式也要有变化（满铺/左侧渐变/右侧竖条）。"""
    images = _fake_images(tmp_path)[:1]
    path = tmp_path / "单图.pptx"
    build_pptx(SECTIONS, path, backgrounds=images)
    deck = Presentation(str(path))

    layouts = []
    for slide in list(deck.slides)[2:]:
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        if not pictures:
            continue
        pic = pictures[0]
        layouts.append((round(pic.left / 914400, 2), round(pic.width / 914400, 2)))
    assert len(set(layouts)) >= 2, f"同一张背景也该逐页换版式：{layouts}"
    assert any(left > 7 for left, _w in layouts), "应该有一页是右侧竖条"


def test_backdrop_count_is_capped(tmp_path):
    images = _fake_images(tmp_path)
    # 同一张图重复传多次只会去重；这里验证 build_pptx 不会因为传入很多而错乱
    path = tmp_path / "多张.pptx"
    build_pptx(SECTIONS, path, backgrounds=[*images, *images, *images, *images])
    assert len(Presentation(str(path)).slides) == 7
    assert MAX_BACKDROPS >= 2
