"""论文结构解析：真实章节、目录跳过、进度推断、工作区清单。"""

from __future__ import annotations

import time

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from paper_agent.core.models import Attachment, ChatSession, Message
from paper_agent.services import paper_outline as po
from paper_agent.services.skills.document_skills import build_docx, build_pdf

TOC_ENTRIES = [
    ("toc 1", "基于YOLOv8的车牌识别检测算法研究\t1"),
    ("toc 1", "目录\t2"),
    ("toc 1", "第一章 绪论\t6"),
    ("toc 2", "1.1 研究背景\t6"),
    ("toc 3", "1.1.1 传统方法\t7"),
    ("toc 1", "第二章 相关工作\t10"),
]


def _write_toc_docx(path, toc_style: bool = True) -> None:
    """还原真实文件的形态：标题 + 摘要 + 目录（含条目）+ 正文各章。"""
    doc = Document()
    if toc_style:
        for name in ("toc 1", "toc 2", "toc 3"):
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    doc.add_heading("基于YOLOv8的车牌识别检测算法研究", level=1)
    doc.add_paragraph("摘要")
    doc.add_paragraph("本文提出了一种改进的车牌检测方法。")
    doc.add_paragraph("关键词：YOLOv8；车牌检测")
    doc.add_heading("目录", level=1)
    for style, text in TOC_ENTRIES:
        if toc_style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)          # 没有 toc 样式时靠「制表符 + 页码」识别

    doc.add_heading("第一章 绪论", level=1)
    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("自动识别车牌在智能交通里非常重要。")
    doc.add_heading("1.1.1 传统方法", level=3)
    doc.add_paragraph("传统方法依赖手工特征。")
    doc.add_heading("第二章 相关工作", level=1)
    doc.add_paragraph("深度学习方法已成为主流。")
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] 张三. 车牌识别研究[J]. 计算机学报, 2025.")
    doc.save(str(path))


# ---------------------------------------------------------------- docx
@pytest.mark.parametrize("toc_style", [True, False])
def test_toc_entries_are_not_chapters(tmp_path, toc_style):
    path = tmp_path / "论文.docx"
    _write_toc_docx(path, toc_style=toc_style)
    nodes = po.extract_outline(path)

    titles = [node.title for node in nodes]
    assert titles == ["摘要", "第一章 绪论", "第二章 相关工作", "参考文献"], titles
    assert "目录" not in titles, "目录是导航信息，不该出现在结构里"
    assert not [t for t in titles if "\t" in t], "目录条目的页码不该混进标题"

    chapter = nodes[1]
    assert [child.title for child in chapter.children] == ["1.1 研究背景"]
    assert chapter.children[0].words > 0
    assert nodes[1].children[0].children[0].title == "1.1.1 传统方法"


def test_document_title_is_dropped(tmp_path):
    path = tmp_path / "论文.docx"
    _write_toc_docx(path)
    titles = [node.title for node in po.extract_outline(path)]
    assert "基于YOLOv8的车牌识别检测算法研究" not in titles


def test_empty_first_chapter_is_kept(tmp_path):
    """不能把「还没写的第一章」误当成文档标题删掉。"""
    path = tmp_path / "空章.docx"
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_heading("第二章 相关工作", level=1)
    doc.add_paragraph("第二章已经写好了。")
    doc.save(str(path))

    nodes = po.extract_outline(path)
    assert [node.title for node in nodes] == ["第一章 绪论", "第二章 相关工作"]
    assert nodes[0].state == "todo" and nodes[1].state == "done"


def test_normal_headings_unaffected(tmp_path):
    path = tmp_path / "正常.docx"
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("1.1 节介绍了背景，这句不是标题。")
    doc.add_heading("第二章 实验（2024）", level=1)
    doc.add_paragraph("2026 年的数据。")
    doc.save(str(path))

    nodes = po.extract_outline(path)
    assert [node.title for node in nodes] == ["第一章 绪论", "第二章 实验（2024）"]
    assert nodes[0].words > 0


def test_only_toc_yields_no_structure(tmp_path):
    path = tmp_path / "只有目录.docx"
    doc = Document()
    doc.add_heading("目录", level=1)
    doc.add_paragraph("第一章 绪论\t6")
    doc.add_paragraph("第二章 相关工作\t10")
    doc.save(str(path))
    assert po.extract_outline(path) == []


def test_progress_state(tmp_path):
    path = tmp_path / "进度.docx"
    build_docx(
        "# 第一章 绪论\n\n## 1.1 研究背景\n\n# 第二章 结论\n\n这里已经有正文了。\n",
        path,
        style="thesis",
    )
    nodes = po.extract_outline(path)
    assert nodes[0].state == "todo" and nodes[0].children[0].state == "todo"
    assert nodes[1].state == "done"


def test_outline_cache_invalidated_on_change(tmp_path):
    path = tmp_path / "缓存.docx"
    build_docx("# 第一章 绪论\n\n正文。\n", path, style="thesis")
    assert [n.title for n in po.extract_outline(path)] == ["第一章 绪论"]

    time.sleep(0.01)
    build_docx("# 第一章 绪论\n\n正文。\n\n# 第二章 相关工作\n\n正文。\n", path, style="thesis")
    titles = [n.title for n in po.extract_outline(path)]
    assert titles == ["第一章 绪论", "第二章 相关工作"], "文件变了应重新解析"


# ---------------------------------------------------------------- pdf
def test_pdf_outline_skips_toc_page(tmp_path):
    markdown = """# 目录

第一章 绪论 .......... 6

第二章 相关工作 .......... 12

# 第一章 绪论

本文研究车牌检测。

# 第二章 相关工作

深度学习方法已成为主流。
"""
    path = tmp_path / "论文.pdf"
    build_pdf(markdown, path)
    titles = [node.title for node in po.extract_outline(path)]
    assert "目录" not in titles
    assert "第一章 绪论" in titles and "第二章 相关工作" in titles


# ---------------------------------------------------------------- 流式预览
def test_markdown_outline_marks_writing_section():
    markdown = "# 第一章 绪论\n\n## 1.1 研究背景\n\n正文内容。"
    nodes = po.extract_markdown_outline(markdown)
    assert [node.title for node in nodes] == ["第一章 绪论"]
    assert [child.title for child in nodes[0].children] == ["1.1 研究背景"]
    assert nodes[0].children[-1].state == "doing"


def test_markdown_outline_skips_code_fence():
    markdown = "# 第一章 绪论\n\n```docx\n# 第二章 不该出现\n```\n"
    nodes = po.extract_markdown_outline(markdown)
    assert [node.title for node in nodes] == ["第一章 绪论"]


def test_markdown_outline_empty_cases():
    assert po.extract_markdown_outline("") == []
    assert po.extract_markdown_outline("没有任何标题的一段话。") == []


def test_outline_signature_covers_children():
    before = po.extract_markdown_outline("# 第一章\n\n## 1.1 小节\n\n正文")
    after = po.extract_markdown_outline("# 第一章\n\n## 1.1 小节\n\n## 1.2 新增\n\n正文")
    assert po.outline_signature(before) != po.outline_signature(after), "子节变化也要能被检测到"


# ---------------------------------------------------------------- 工作区清单
def _session_with_files(tmp_path):
    first = tmp_path / "第一篇.docx"
    second = tmp_path / "第二篇.pdf"
    first.write_bytes(b"a")
    second.write_bytes(b"bb")
    session = ChatSession(title="多篇")
    now = time.time()
    session.add_message(
        Message(
            role="assistant",
            content="",
            artifacts=[
                Attachment(name=first.name, path=str(first), kind="file", created_at=now - 60),
                Attachment(name=second.name, path=str(second), kind="file", created_at=now),
            ],
        )
    )
    return session, first, second


def test_paper_candidates_sorted_by_time(tmp_path):
    session, first, second = _session_with_files(tmp_path)
    candidates = po.paper_candidates(session)
    assert [item.name for item in candidates] == [second.name, first.name]
    assert po.source_label(candidates[0]).endswith("模型生成")


def test_workspace_snapshot_fields(tmp_path):
    session, first, _second = _session_with_files(tmp_path)
    snapshot = po.workspace_snapshot(session)
    assert len(snapshot) == 2
    assert {item["name"] for item in snapshot} == {"第一篇.docx", "第二篇.pdf"}
    assert all(item["path"] and item["source"] == "model" for item in snapshot)

    line = po.format_workspace_file(snapshot[0])
    assert snapshot[0]["name"] in line and snapshot[0]["path"] in line


def test_extract_outline_invalid_inputs(tmp_path):
    assert po.extract_outline(tmp_path / "不存在.docx") == []
    unknown = tmp_path / "笔记.txt"
    unknown.write_text("内容", encoding="utf-8")
    assert po.extract_outline(unknown) == []
