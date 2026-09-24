"""文档 Skills：论文排版、配图、目录域、三线表、题注、读取与产物路径。"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from paper_agent.services.skills import document_skills as ds

THESIS_MD = """# 基于YOLOv8的车牌识别检测算法研究

# 摘要

本文提出了一种改进的车牌检测方法。

关键词：YOLOv8；车牌检测

# 目录

第一章 绪论\t6

1.1 研究背景\t6

# 第一章 绪论

## 1.1 研究背景

自动识别车牌在智能交通里非常重要。

表 3-1 实验环境配置

| 项目 | 配置 |
| --- | --- |
| CPU | i7-12700 |
| GPU | RTX 4090 |

图 3-1 系统整体架构

![系统整体架构图]({image})

# 参考文献

[1] 张三. 车牌识别研究[J]. 计算机学报, 2025.
"""


def _xpath(element, path: str):
    return element.findall(path, {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})


def _build(tmp_path, markdown: str, name: str = "毕业论文.docx", style: str = "thesis"):
    path = tmp_path / name
    ds.build_docx(markdown, path, style=style)
    return path, Document(str(path))


# ---------------------------------------------------------------- 排版基础
def test_thesis_typography(tmp_path):
    _, doc = _build(tmp_path, "# 第一章 绪论\n\n正文内容。\n")
    normal = doc.styles["Normal"]
    assert normal.font.size.pt == 12                                   # 小四
    assert normal.element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert normal.paragraph_format.first_line_indent.pt == 24          # 首行缩进 2 字符
    assert abs(normal.paragraph_format.line_spacing - 1.5) < 1e-6

    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"
        assert style.font.size.pt == size
        assert style.font.bold is True
        assert str(style.font.color.rgb) == "000000"


def test_plain_style_unchanged(tmp_path):
    """普通文档不套论文规范（用户明确要求过）。"""
    _, doc = _build(tmp_path, "# 会议记录\n\n讨论排期。\n", "会议记录.docx", style="auto")
    normal = doc.styles["Normal"]
    assert normal.font.name == "Microsoft YaHei"
    assert normal.font.size.pt == 11
    assert normal.paragraph_format.line_spacing is None


def test_doc_type_detection():
    assert ds.resolve_doc_style("auto", "# 第一章 绪论\n\n正文", "a.docx") == "thesis"
    assert ds.resolve_doc_style("auto", "# 会议记录\n\n排期", "b.docx") == "plain"
    assert ds.resolve_doc_style("thesis", "# 会议记录", "b.docx") == "thesis"
    assert ds.resolve_doc_style("plain", "# 第一章 绪论", "a.docx") == "plain"


# ---------------------------------------------------------------- 配图
def test_image_block_parsed(tmp_path, sample_png):
    blocks = ds.parse_markdown_blocks(f"![系统架构图]({sample_png})")
    assert blocks == [("image", f"系统架构图\n{sample_png}")]
    assert ds._split_image(blocks[0][1]) == ("系统架构图", str(sample_png))


def test_image_inline_still_stripped():
    """行内图片（夹在文字里）仍按普通文本处理，不生成图片块。"""
    blocks = ds.parse_markdown_blocks("结果见 ![图](a.png) 所示。")
    assert [kind for kind, _ in blocks] == ["p"]
    assert "a.png" not in blocks[0][1]


def test_image_inserted_with_caption(tmp_path, sample_png):
    markdown = f"# 第一章 绪论\n\n图 3-1 系统整体架构\n\n![系统整体架构图]({sample_png})\n"
    _, doc = _build(tmp_path, markdown)
    assert len(doc.inline_shapes) == 1, "配图应真的插进 Word"
    texts = [p.text for p in doc.paragraphs]
    assert "图 3-1 系统整体架构" in texts
    assert texts.count("图 3-1 系统整体架构") == 1, "题注不应重复（alt 与下一行取一个）"
    caption = next(p for p in doc.paragraphs if p.text.startswith("图 3-1"))
    assert caption.text == "图 3-1 系统整体架构"
    assert not caption.paragraph_format.first_line_indent, "题注不应带首行缩进"


def test_image_scaled_to_page_width(tmp_path, sample_png):
    _, doc = _build(tmp_path, f"# 第一章\n\n![超大图]({sample_png})\n")
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    assert doc.inline_shapes[0].width <= usable


def test_missing_image_placeholder(tmp_path):
    _, doc = _build(tmp_path, "# 第一章\n\n![缺失的图](不存在的图.png)\n")
    assert not doc.inline_shapes
    assert any("图片未找到" in p.text for p in doc.paragraphs)


def test_image_found_by_name_in_artifacts(tmp_path, sample_png, artifacts_dir):
    """模型常只给文件名：应能在产物目录里找到。"""
    import shutil

    shutil.copy(sample_png, artifacts_dir / "配图-1.png")
    _, doc = _build(tmp_path, "# 第一章\n\n![架构图](配图-1.png)\n")
    assert len(doc.inline_shapes) == 1


# ---------------------------------------------------------------- 目录域
def test_toc_field_replaces_manual_toc(tmp_path, sample_png):
    path, doc = _build(tmp_path, THESIS_MD.format(image=sample_png))
    instr = [el.text for el in doc.element.body.iter(qn("w:instrText"))]
    assert any("TOC" in (text or "") for text in instr), "应插入 Word 目录域"

    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "1.1 研究背景\t6" not in body_text, "手写目录行应被目录域取代"
    assert "目录" in body_text
    # 目录标题不能是 Heading 1，否则目录会把自己也列进去
    toc_heading = next(p for p in doc.paragraphs if p.text.strip() == "目录")
    assert toc_heading.style.name != "Heading 1"


def test_update_fields_enabled(tmp_path):
    _, doc = _build(tmp_path, THESIS_MD.format(image="x.png"))
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_page_number_footer(tmp_path):
    _, doc = _build(tmp_path, THESIS_MD.format(image="x.png"))
    footer = doc.sections[0].footer
    instr = [el.text for el in footer._element.iter(qn("w:instrText"))]
    assert any("PAGE" in (text or "") for text in instr), "页脚应有页码域"


def test_plain_doc_has_no_toc_or_footer(tmp_path):
    _, doc = _build(tmp_path, "# 目录\n\n第一章 绪论\t6\n", "记录.docx", style="plain")
    assert not any("TOC" in (el.text or "") for el in doc.element.body.iter(qn("w:instrText")))
    assert not any(
        "PAGE" in (el.text or "") for el in doc.sections[0].footer._element.iter(qn("w:instrText"))
    )


# ---------------------------------------------------------------- 三线表与题注
def _table_borders(table):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    return borders


def test_three_line_table(tmp_path):
    _, doc = _build(tmp_path, THESIS_MD.format(image="x.png"))
    table = doc.tables[0]
    borders = _table_borders(table)
    assert borders is not None
    values = {child.tag.split("}")[1]: child.get(qn("w:val")) for child in borders}
    assert values["top"] == "single" and values["bottom"] == "single"   # 上下框线
    assert values["insideV"] == "none"                                 # 无竖线
    assert values["insideH"] == "none"                                 # 无内部横线
    assert values["left"] == "none" and values["right"] == "none"

    header_bottom = table.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    assert header_bottom is not None, "表头行下方应有分隔线"
    assert header_bottom.find(qn("w:bottom")).get(qn("w:val")) == "single"


def test_table_header_repeats_across_pages(tmp_path):
    _, doc = _build(tmp_path, THESIS_MD.format(image="x.png"))
    assert doc.tables[0].rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None


def test_table_caption_above(tmp_path):
    _, doc = _build(tmp_path, THESIS_MD.format(image="x.png"))
    texts = [p.text for p in doc.paragraphs]
    assert "表 3-1 实验环境配置" in texts
    caption_index = texts.index("表 3-1 实验环境配置")
    # 题注后紧跟表格：正文里题注之后第一个非空段落应是表格之后的内容
    assert texts[caption_index + 1].strip() in {"", "图 3-1 系统整体架构"}


def test_plain_table_keeps_previous_behaviour(tmp_path):
    _, doc = _build(tmp_path, "| a | b |\n| --- | --- |\n| 1 | 2 |\n", "t.docx", style="plain")
    assert _table_borders(doc.tables[0]) is None, "普通文档表格不加三线表格式"


# ---------------------------------------------------------------- 读取与产物路径
def test_read_document_reads_tables(tmp_path):
    path, _ = _build(tmp_path, THESIS_MD.format(image="x.png"))
    text = ds.read_document(str(path))
    assert "i7-12700" in text and "RTX 4090" in text, "读文档时要能看到表格内容"
    assert "项目" in text


def test_unique_path_overwrites_same_session_file(tmp_path, artifacts_dir):
    """同一会话里重新生成同名文档 → 原地更新，不再堆 -1/-2 副本。"""
    first = ds._unique_path("论文.docx")
    first.write_bytes(b"v1")
    assert ds._unique_path("论文.docx", {str(first)}) == first
    other = ds._unique_path("论文.docx", set())
    assert other.name == "论文-1.docx"


def test_unique_path_ignores_foreign_session(tmp_path, artifacts_dir):
    first = ds._unique_path("论文.docx")
    first.write_bytes(b"v1")
    second = ds._unique_path("论文.docx", {str(artifacts_dir / "别的文件.docx")})
    assert second.name == "论文-1.docx"


def test_create_docx_tool_reports_style(tmp_path, artifacts_dir):
    plain = ds.CreateDocxTool().run(filename="会议记录.docx", markdown="# 会议记录\n\n排期")
    assert "普通文档排版" in plain.content
    paper = ds.CreateDocxTool().run(filename="毕业论文.docx", markdown="# 第一章 绪论\n\n正文")
    assert "论文排版" in paper.content


@pytest.mark.parametrize("name", ["报告", "报告.docx", "报告.pdf"])
def test_create_docx_tool_normalises_name(tmp_path, artifacts_dir, name):
    result = ds.CreateDocxTool().run(filename=name, markdown="# 标题\n\n正文")
    assert result.artifact_paths[0].endswith(".docx")


# ---------------------------------------------------------------- 端到端
def test_generated_thesis_parses_back(tmp_path, sample_png):
    """自己生成的论文要能被结构化解析回章节树（两个模块的约定必须一致）。"""
    from paper_agent.services import paper_outline as po

    path, _doc = _build(tmp_path, THESIS_MD.format(image=sample_png))
    nodes = po.extract_outline(path)
    titles = [node.title for node in nodes]

    assert "摘要" in titles and "第一章 绪论" in titles and "参考文献" in titles
    assert "目录" not in titles, "目录域不该被当成章节"
    assert "基于YOLOv8的车牌识别检测算法研究" not in titles, "文档标题不该算一章"

    chapter = next(node for node in nodes if node.title == "第一章 绪论")
    assert [child.title for child in chapter.children] == ["1.1 研究背景"]
    assert next(node for node in nodes if node.title == "摘要").words > 0
