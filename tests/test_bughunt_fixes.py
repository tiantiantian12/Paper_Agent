"""缺陷审计中确认的问题的回归用例。"""

from __future__ import annotations

import re
import time

import pytest
from pptx import Presentation
from pptx.util import Inches

from paper_agent.core.models import Attachment, ChatSession, Message, OutlineNode
from paper_agent.services.paper_outline import extract_outline, outline_signature
from paper_agent.services.reasoning import is_unsupported_error, resolve_mode
from paper_agent.services.skills.document_skills import build_pptx


# ---------------------------------------------------------------- PPT：封面不丢内容
def _slide_texts(slide) -> list[str]:
    return [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


def _all_text(deck) -> str:
    return "\n".join(text for slide in deck.slides for text in _slide_texts(slide))


def _build(tmp_path, markdown: str, name: str = "答辩.pptx"):
    path = tmp_path / name
    build_pptx(markdown, path)
    return Presentation(str(path))


def test_cover_keeps_bullets_of_similar_first_section(tmp_path):
    """首节标题与文件名相近、且带要点时，不能把这一页内容吞掉。"""
    deck = _build(
        tmp_path,
        "# 研究背景与意义\n\n- 要点甲\n- 要点乙\n- 要点丙\n\n"
        "# 国内外研究现状\n\n- 现状\n\n# 核心算法\n\n- 算法\n\n"
        "# 实验结果\n\n- 结果\n\n# 结论与展望\n\n- 结论\n",
        name="研究背景与意义_v3.pptx",
    )
    text = _all_text(deck)
    assert all(item in text for item in ("要点甲", "要点乙", "要点丙")), text
    assert "研究背景与意义" in _slide_texts(deck.slides[0])[0]        # 标题借给封面用


def test_similar_first_section_is_not_listed_in_toc(tmp_path):
    """借去当标题的那一节不进目录（否则封面标题会变成目录第一条）。"""
    deck = _build(
        tmp_path,
        "# 研究背景与意义\n\n- 要点甲\n\n"
        "# 一、现状\n\n- a\n\n# 二、算法\n\n- b\n\n"
        "# 三、实验\n\n- c\n\n# 四、部署\n\n- d\n\n# 五、结论\n\n- e\n",
        name="研究背景与意义.pptx",
    )
    toc = next(
        slide for slide in deck.slides if _slide_texts(slide)[0].strip() == "目录"
    )
    entries = _toc_entries(toc)
    assert len(entries) == 5, entries
    assert not any("研究背景与意义" in entry for entry in entries)


def test_cover_keeps_long_paragraph_of_first_section(tmp_path):
    """首节只有长段落（不构成封面短句）时，整节必须保留为正文页。"""
    long_text = "本文围绕车牌识别展开研究。" * 12          # > 60 字
    deck = _build(
        tmp_path,
        f"# 研究背景\n\n{long_text}\n\n# 第一章 绪论\n\n- x\n",
        name="车牌识别答辩.pptx",
    )
    assert long_text[:20] in _all_text(deck), "长段落不能被当作封面副标题吞掉"
    assert len(deck.slides) >= 3


def test_cover_still_consumes_pure_title_section(tmp_path):
    """「标题 + 一句短话」仍然整节当封面（原有行为不变）。"""
    deck = _build(tmp_path, "# 开题答辩\n\n汇报人：张三\n\n# 一、背景\n\n- a\n")
    assert _slide_texts(deck.slides[0])[0].startswith("开题答辩")
    assert "汇报人：张三" in _all_text(deck)


# ---------------------------------------------------------------- PPT：目录不显示页码
def _toc_entries(slide) -> list[str]:
    entries: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        if shape.top is None or shape.top > Inches(6.5):      # 页脚（讲题名 / 页码）
            continue
        stripped = shape.text_frame.text.strip()
        if stripped == "目录" or stripped.startswith("（还有"):
            continue
        entries.append(stripped)
    return entries


def test_toc_entries_have_no_page_number_prefix(tmp_path):
    """目录条目不显示页码前缀（「02 一、研究背景与意义」这种写法已去掉）。"""
    bullets = "\n".join(f"- 要点{i}" for i in range(1, 13))        # 拆成 2 页
    deck = _build(
        tmp_path,
        f"# 讲题\n\n# 一、背景\n\n- a\n\n# 二、核心算法\n\n{bullets}\n\n"
        "# 三、实验\n\n- c\n\n# 四、部署\n\n- d\n\n# 五、结论\n\n- e\n",
        name="答辩.pptx",
    )
    toc = next(
        slide for slide in deck.slides if _slide_texts(slide)[0].strip() == "目录"
    )
    entries = _toc_entries(toc)

    assert entries == ["一、背景", "二、核心算法", "三、实验", "四、部署", "五、结论"]
    assert not any(re.match(r"^\d", entry) for entry in entries)


# ---------------------------------------------------------------- PPT：表格截断有提示
def test_oversized_table_is_not_silently_cut(tmp_path):
    rows = "| 模型 | mAP | 召回 |\n| --- | --- | --- |\n" + "\n".join(
        f"| 模型{i} | 0.{i} | 0.{i} |" for i in range(1, 13)
    )
    deck = _build(
        tmp_path,
        f"# 讲题\n\n# 对比实验\n\n{rows}\n\n# 结论\n\n- c\n",
        name="对比.pptx",
    )
    text = _all_text(deck)
    assert "只显示前" in text, "被截断的表格必须写明，不能静默丢行"
    assert "13" in text or "12" in text


# ---------------------------------------------------------------- 结构指纹含状态
def test_signature_catches_state_change():
    before = [OutlineNode("第一章", 10, "doing"), OutlineNode("第二章", 5, "todo")]
    after = [OutlineNode("第一章", 10, "done"), OutlineNode("第二章", 5, "todo")]
    assert outline_signature(before) != outline_signature(after)


def test_parse_failure_is_not_cached(tmp_path, monkeypatch):
    """解析失败（文件正被写入等瞬时原因）不该把空结果缓存下来。"""
    from paper_agent.services import paper_outline

    doc = tmp_path / "论文.docx"
    doc.write_bytes(b"not a real docx")
    calls = {"n": 0}

    def boom(_path):
        calls["n"] += 1
        raise ValueError("文件被占用")

    monkeypatch.setattr(paper_outline, "_docx_outline", boom)
    assert extract_outline(doc) == []
    assert extract_outline(doc) == []
    assert calls["n"] == 2, "失败结果不该进缓存"


# ---------------------------------------------------------------- 推理参数误判
@pytest.mark.parametrize("model_id", ["cogito3", "gpt-4o-mini", "qwen-max", "abc123"])
def test_auto_mode_ignores_lookalike_names(model_id):
    assert resolve_mode("auto", model_id) == "off"


@pytest.mark.parametrize("model_id", ["o3-mini", "gpt-5", "o1", "qwq-32b"])
def test_auto_mode_still_detects_real_ones(model_id):
    assert resolve_mode("auto", model_id) == "reasoning_effort"


@pytest.mark.parametrize(
    "text",
    [
        "unknown parameter: temperature",
        "model not supported",
        "invalid parameter: max_tokens",
    ],
)
def test_unsupported_error_needs_reasoning_context(text):
    """与推理无关的 400 不能触发「撤掉推理参数」。"""
    assert is_unsupported_error(text) is False


@pytest.mark.parametrize(
    "text",
    [
        "reasoning_effort is not supported",
        "不支持推理参数：reasoning_effort",
    ],
)
def test_unsupported_error_detects_reasoning_context(text):
    assert is_unsupported_error(text) is True


# ---------------------------------------------------------------- 自动重试与面板
def _window(qapp, tmp_path, sessions):
    from paper_agent.core.config import AppConfig
    from paper_agent.core.session_store import SessionStore
    from paper_agent.ui.main_window import MainWindow

    window = MainWindow(AppConfig(), SessionStore(tmp_path / "sessions"))
    window.sessions = {session.id: session for session in sessions}
    window.current = sessions[0] if sessions else None
    return window


def _failing_session() -> tuple[ChatSession, Message, Message]:
    session = ChatSession(title="会话")
    user = Message(role="user", content="写点东西")
    assistant = Message(role="assistant", content="写了一半", status="streaming")
    session.add_message(user)
    session.add_message(assistant)
    return session, user, assistant


def test_pending_retry_is_cancelled_on_new_messages(qapp, tmp_path):
    """用户在等待期间又发消息时，必须作废重试（否则重试会删掉这些消息）。"""
    session, _user, assistant = _failing_session()
    window = _window(qapp, tmp_path, [session])
    window.chat_view.add_message(assistant)

    window._on_agent_failed(assistant.id, "HTTP 429 · tpm limit")
    assert window._retry_timer.isActive(), "限流应排定一次较长的重试"
    assert window._retry_message_id == assistant.id

    window._cancel_pending_retry()
    assert not window._retry_timer.isActive()
    assert window._retry_message_id == ""
    assert "自动重试" not in window.chat_view.widget_for(assistant.id).status_label.text()


def test_pending_retry_is_cancelled_on_session_switch(qapp, tmp_path):
    first, _user, assistant = _failing_session()
    second = ChatSession(title="另一个")
    window = _window(qapp, tmp_path, [first, second])

    window._on_agent_failed(assistant.id, "HTTP 500 · server error")
    assert window._retry_timer.isActive()

    window.activate_session(second.id)
    assert not window._retry_timer.isActive(), "切走后重试不该再打到旧会话"
    assert window._auto_retry_count == 0


def test_pending_retry_refuses_when_conversation_moved_on(qapp, tmp_path, monkeypatch):
    """失败那条后面又有了新消息 → 放弃重试，绝不删除后续对话。"""
    session, _user, assistant = _failing_session()
    later = Message(role="user", content="顺便再改一下")
    session.add_message(later)
    window = _window(qapp, tmp_path, [session])

    called = {"n": 0}
    monkeypatch.setattr(
        window, "_auto_retry", lambda message_id: called.__setitem__("n", called["n"] + 1)
    )
    window._retry_message_id = assistant.id
    window._run_pending_retry()

    assert called["n"] == 0, "会话已往后走，不能重试"
    assert [m.id for m in session.messages][-1] == later.id
    assert len(session.messages) == 3


def test_tick_outline_leaves_other_session_alone(qapp, tmp_path, monkeypatch):
    """后台会话在生成时，切过去的会话结构面板不该被清空。"""
    from paper_agent.services.agent_service import AgentService
    from paper_agent.services.skills.document_skills import build_docx

    paper = tmp_path / "论文.docx"
    build_docx("# 第一章 绪论\n\n正文内容。\n", paper, style="thesis")

    streaming = ChatSession(title="正在生成")
    assistant = Message(role="assistant", content="写入中", status="streaming")
    streaming.add_message(assistant)

    other = ChatSession(title="另一个")
    other.add_message(
        Message(
            role="assistant",
            content="",
            artifacts=[
                Attachment(
                    name=paper.name, path=str(paper), kind="file", created_at=time.time()
                )
            ],
        )
    )

    window = _window(qapp, tmp_path, [streaming, other])
    window.show()
    qapp.processEvents()
    window.activate_session(other.id)
    assert window.outline_panel.tree.topLevelItemCount() > 0

    real = AgentService.is_running
    AgentService.is_running = property(lambda self: True)
    try:
        window._streaming_message_id = assistant.id      # 流式消息属于另一个会话
        window.outline_panel.setVisible(True)
        window._outline_signature = None
        window._tick_outline()
    finally:
        AgentService.is_running = real

    assert window.outline_panel.tree.topLevelItemCount() > 0, "不该把当前会话的面板画空"
